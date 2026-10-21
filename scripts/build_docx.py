#!/usr/bin/env python3
"""build_docx.py - turn a manuscript written in Markdown into a Word file.

Arabic paragraphs are set right-to-left automatically (bidi paragraph, RTL runs,
complex-script font), so a bilingual thesis renders correctly in Word.

Supported Markdown: # headings (H1 = centred title), paragraphs, **bold**,
*italic*, - bullets, 1. numbered lists, pipe tables, > block quotes. Write
significance stars as \\* (e.g. .31\\*\\*). <!-- comments --> are dropped.
Paragraphs under a 'References' / 'المراجع' heading get a hanging indent (APA).
Tables default to APA style (horizontal rules only); --table-style grid for all borders.

University presets (--preset NAME, from references/university-presets.json):
  margins with a binding side, fonts, sizes, spacing, heading sizes, page-number
  position, front-matter numbering (roman / abjad letters) and body numbering
  from 1 via automatic section breaks, a cover page, a Table of Contents field
  and lists of tables / figures. --list-presets shows what exists and whether
  each preset was verified against the university's own guide.

Cover page: put a YAML-ish block at the top of the Markdown:
  ---
  title: ...        author: ...      supervisor: ...   co_supervisor: ...
  university: ...   faculty: ...     department: ...   degree: ...
  program: ...      year: ...        date: ...         logo: path.png
  ---
The preset's cover list decides which fields appear and in what order (AR or EN
labels are chosen from the title's script, or --lang). An optional line
<!-- body --> marks where chapter 1 starts (body page numbering restarts at 1);
without it the first ## heading that is not front matter (abstract, dedication,
acknowledgements, abbreviations) starts the body.

Word fills the TOC / lists itself: open the file and accept the "update fields"
prompt, or press Ctrl+A then F9. python-docx cannot compute page numbers.

Usage:
  python -X utf8 build_docx.py paper.md -o paper.docx
  python -X utf8 build_docx.py thesis_ar.md -o thesis.docx --lang ar --spacing 1.5
  python -X utf8 build_docx.py thesis.md -o thesis.docx --preset generic-arab
Options: --font, --ar-font, --size, --ar-size, --spacing, --margin-cm, --margins-cm T,B,L,R,
         --mirror-margins, --preset, --list-presets
"""

import argparse
import json
import os
import re
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

try:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    sys.exit("python-docx is missing: pip install python-docx")

ARABIC_RE = re.compile(r"[؀-ۿ]")
REF_HEADINGS = re.compile(r"^(references|bibliography|reference list|works cited|المراجع|قائمة المراجع|المصادر والمراجع)$", re.I)
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)")
PRESETS_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "references", "university-presets.json")


def is_arabic(text, lang):
    if lang in ("ar", "en"):
        return lang == "ar"
    letters = re.findall(r"[A-Za-z؀-ۿ]", text)
    return bool(letters) and sum(1 for c in letters if ARABIC_RE.match(c)) / len(letters) > 0.4


def set_rtl_paragraph(p):
    ppr = p._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    ppr.append(bidi)


def style_run(run, opts, rtl, bold=False, italic=False, mono=False):
    run.bold = bold or None
    run.italic = italic or None
    font = "Consolas" if mono else (opts.ar_font if rtl else opts.font)
    size = opts.ar_size if rtl else opts.size
    run.font.name = font
    run.font.size = Pt(size)
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), opts.ar_font if attr == "w:cs" else font)
    szcs = OxmlElement("w:szCs")
    szcs.set(qn("w:val"), str(int(opts.ar_size * 2)))
    rpr.append(szcs)
    if rtl:
        rpr.append(OxmlElement("w:rtl"))
        if bold:
            rpr.append(OxmlElement("w:bCs"))
        if italic:
            rpr.append(OxmlElement("w:iCs"))


STAR = "\x01"      # stands in for an escaped \* (significance stars) while parsing


# Inside an Arabic paragraph, numbers, page ranges, DOIs and Latin words go in
# their own left-to-right runs: in an RTL run Word reorders '76-107' to
# '107-76' and breaks ranges across lines.
LTR_RUN = re.compile(r"(https?://\S+|10\.\d{4,9}/\S+|\.?[A-Za-z0-9][A-Za-z0-9.\-/:_%]*[A-Za-z0-9%]|\.?[A-Za-z0-9])")


def put_runs(p, text, opts, rtl, **fmt):
    if not rtl:
        style_run(p.add_run(text), opts, False, **fmt)
        return
    for k, seg in enumerate(LTR_RUN.split(text)):
        if seg:
            is_ltr = k % 2 == 1
            # A non-breaking hyphen keeps '76-107' on one line.
            run = p.add_run(seg.replace("-", "‑") if is_ltr and re.fullmatch(r"[\d.\-]+", seg) else seg)
            style_run(run, opts, not is_ltr, **fmt)
            if is_ltr:
                run.font.size = Pt(opts.ar_size)            # same line height as the Arabic around it
                if not re.search(r"[A-Za-z]", seg) and not fmt.get("mono"):
                    run.font.name = opts.ar_font            # digits match the Arabic text


def add_inline(p, text, opts, rtl, base_bold=False):
    text = text.replace("\\*", STAR)
    put = lambda s: s.replace(STAR, "*")
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            put_runs(p, put(tok[2:-2]), opts, rtl, bold=True)
        elif tok.startswith("`") and tok.endswith("`"):
            put_runs(p, put(tok[1:-1]), opts, rtl, mono=True)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            put_runs(p, put(tok[1:-1]), opts, rtl, bold=base_bold, italic=True)
        else:
            put_runs(p, put(tok), opts, rtl, bold=base_bold)


def line_spacing(opts, rtl, hanging=False):
    if hanging and getattr(opts, "refs_spacing", None):
        return opts.refs_spacing
    if rtl and getattr(opts, "ar_spacing", None):
        return opts.ar_spacing
    return opts.spacing


def fmt_paragraph(p, opts, rtl, hanging=False, first_line=False):
    pf = p.paragraph_format
    pf.line_spacing = line_spacing(opts, rtl, hanging)
    pf.space_after = Pt(6)
    if rtl:
        set_rtl_paragraph(p)
    # In a bidi paragraph Word reads LEFT as the start edge (right side), so
    # LEFT = ragged at the end in both directions. References use it (APA wants
    # them unjustified, and justified Arabic lines with long DOIs stretch badly);
    # body text is justified unless the preset says otherwise (APA 7, KSU).
    justify = getattr(opts, "justify", True)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if (hanging or not justify) else WD_ALIGN_PARAGRAPH.JUSTIFY
    if hanging:
        pf.left_indent = Cm(1.27)
        pf.first_line_indent = Cm(-1.27)
    elif first_line:
        pf.first_line_indent = Cm(1.27)


def add_field(run, instr):
    """Append a complex field (fldChar begin / instrText / fldChar end) to a run."""
    for tag, text in (("begin", None), (None, instr), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


PAGE_POS = {"bottom-center": ("footer", WD_ALIGN_PARAGRAPH.CENTER), "bottom-right": ("footer", WD_ALIGN_PARAGRAPH.RIGHT),
            "bottom-left": ("footer", WD_ALIGN_PARAGRAPH.LEFT), "top-center": ("header", WD_ALIGN_PARAGRAPH.CENTER),
            "top-right": ("header", WD_ALIGN_PARAGRAPH.RIGHT), "top-left": ("header", WD_ALIGN_PARAGRAPH.LEFT)}


def add_page_number(section, position="bottom-center"):
    where, align = PAGE_POS.get(position, PAGE_POS["bottom-center"])
    p = (section.footer if where == "footer" else section.header).paragraphs[0]
    p.alignment = align
    add_field(p.add_run(), "PAGE")


# ---- schema-ordered XML inserts (Word rejects children out of order) ----

SECT_ORDER = ["w:headerReference", "w:footerReference", "w:footnotePr", "w:endnotePr", "w:type", "w:pgSz", "w:pgMar",
              "w:paperSrc", "w:pgBorders", "w:lnNumType", "w:pgNumType", "w:cols", "w:formProt", "w:vAlign",
              "w:noEndnote", "w:titlePg", "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid",
              "w:printerSettings", "w:sectPrChange"]
SETTINGS_ORDER = ["w:writeProtection", "w:view", "w:zoom", "w:removePersonalInformation", "w:removeDateAndTime",
                  "w:doNotDisplayPageBoundaries", "w:displayBackgroundShape", "w:printPostScriptOverText",
                  "w:printFractionalCharacterWidth", "w:printFormsData", "w:embedTrueTypeFonts", "w:embedSystemFonts",
                  "w:saveSubsetFonts", "w:saveFormsData", "w:mirrorMargins", "w:alignBordersAndEdges",
                  "w:bordersDoNotSurroundHeader", "w:bordersDoNotSurroundFooter", "w:gutterAtTop",
                  "w:hideSpellingErrors", "w:hideGrammaticalErrors", "w:activeWritingStyle", "w:proofState",
                  "w:formsDesign", "w:attachedTemplate", "w:linkStyles", "w:stylePaneFormatFilter",
                  "w:stylePaneSortMethod", "w:documentType", "w:mailMerge", "w:revisionView", "w:trackRevisions",
                  "w:doNotTrackMoves", "w:doNotTrackFormatting", "w:documentProtection", "w:autoFormatOverride",
                  "w:styleLockTheme", "w:styleLockQFSet", "w:defaultTabStop", "w:autoHyphenation",
                  "w:consecutiveHyphenLimit", "w:hyphenationZone", "w:doNotHyphenateCaps", "w:showEnvelope",
                  "w:summaryLength", "w:clickAndTypeStyle", "w:defaultTableStyle", "w:evenAndOddHeaders",
                  "w:bookFoldRevPrinting", "w:bookFoldPrinting", "w:bookFoldPrintingSheets",
                  "w:drawingGridHorizontalSpacing", "w:drawingGridVerticalSpacing",
                  "w:displayHorizontalDrawingGridEvery", "w:displayVerticalDrawingGridEvery",
                  "w:doNotUseMarginsForDrawingGridOrigin", "w:drawingGridHorizontalOrigin",
                  "w:drawingGridVerticalOrigin", "w:doNotShadeFormData", "w:noPunctuationKerning",
                  "w:characterSpacingControl", "w:printTwoOnOne", "w:strictFirstAndLastChars",
                  "w:noLineBreaksAfter", "w:noLineBreaksBefore", "w:savePreviewPicture",
                  "w:doNotValidateAgainstSchema", "w:saveInvalidXml", "w:ignoreMixedContent",
                  "w:alwaysShowPlaceholderText", "w:doNotDemarcateInvalidXml", "w:saveXmlDataOnly",
                  "w:useXSLTWhenSaving", "w:saveThroughXslt", "w:showXMLTags", "w:alwaysMergeEmptyNamespace",
                  "w:updateFields", "w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat", "w:docVars",
                  "w:rsids", "w:mathPr", "w:attachedSchema", "w:themeFontLang", "w:clrSchemeMapping",
                  "w:doNotIncludeSubdocsInStats", "w:doNotAutoCompressPictures", "w:forceUpgrade", "w:captions",
                  "w:readModeInkLockDown", "w:smartTagType", "w:schemaLibrary", "w:shapeDefaults",
                  "w:doNotEmbedSmartTags", "w:decimalSymbol", "w:listSeparator"]


def insert_ordered(parent, el, order):
    """Replace an existing element of the same tag, else insert before the first later sibling."""
    tag = el.tag
    for old in parent.findall(tag):
        parent.remove(old)
    names = [qn(t) for t in order]
    pos = names.index(tag) if tag in names else len(names)
    for child in parent:
        if child.tag in names and names.index(child.tag) > pos:
            child.addprevious(el)
            return el
    parent.append(el)
    return el


def set_setting(doc, tag, val="true"):
    el = OxmlElement(tag)
    el.set(qn("w:val"), val)
    insert_ordered(doc.settings.element, el, SETTINGS_ORDER)


NUM_FMT = {"roman": "lowerRoman", "ROMAN": "upperRoman", "abjad": "arabicAbjad", "letters": "arabicAbjad",
           "alpha": "arabicAlpha", "decimal": "decimal"}


def set_page_numbering(section, fmt=None, start=1):
    """fmt: a key of NUM_FMT or a Word w:fmt value; None removes the override (plain 1, 2, 3)."""
    sect = section._sectPr
    for old in sect.findall(qn("w:pgNumType")):
        sect.remove(old)
    if fmt is None and start is None:
        return
    el = OxmlElement("w:pgNumType")
    if fmt:
        el.set(qn("w:fmt"), NUM_FMT.get(fmt, fmt))
    if start is not None:
        el.set(qn("w:start"), str(start))
    insert_ordered(sect, el, SECT_ORDER)


def _border(parent, tag, edges):
    box = OxmlElement(tag)
    for edge, on in edges.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single" if on else "nil")
        if on:
            el.set(qn("w:sz"), "8")
            el.set(qn("w:color"), "000000")
        box.append(el)
    parent.append(box)


def tblpr_insert(tblpr, el, before=("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")):
    """Insert in schema order (Word rejects out-of-order table properties)."""
    for tag in before:
        found = tblpr.find(qn(tag))
        if found is not None:
            found.addprevious(el)
            return
    tblpr.append(el)


def apa_borders(t):
    """APA 7 table: horizontal rules above and below the header and at the bottom only."""
    holder = OxmlElement("w:tmp")
    _border(holder, "w:tblBorders", {"top": True, "left": False, "bottom": True, "right": False,
                                     "insideH": False, "insideV": False})
    tblpr_insert(t._tbl.tblPr, holder[0])
    for cell in t.rows[0].cells:
        _border(cell._tc.get_or_add_tcPr(), "w:tcBorders", {"bottom": True})


def add_table(doc, rows, opts, rtl):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
    if not cells:
        return
    ncol = max(len(r) for r in cells)
    t = doc.add_table(rows=len(cells), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if opts.table_style == "grid":
        t.style = "Table Grid"
    else:
        apa_borders(t)
    if rtl:
        tblpr_insert(t._tbl.tblPr, OxmlElement("w:bidiVisual"),
                     before=("w:tblStyleRowBandSize", "w:tblStyleColBandSize", "w:tblW", "w:jc",
                             "w:tblBorders", "w:tblLook"))
    # Column widths follow the longest cell, so '[-0.20, -0.04]' does not wrap mid-number.
    longest = [max(len(r[j]) if j < len(r) else 0 for r in cells) + 2 for j in range(ncol)]
    width = opts.text_width_cm
    for j, n in enumerate(longest):
        w = Cm(width * n / sum(longest))
        for row in t.rows:
            row.cells[j].width = w
    keep = len(cells) <= 25
    for i, r in enumerate(cells):
        trpr = t.rows[i]._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            trpr.append(OxmlElement("w:tblHeader"))      # header repeats on a new page
        for j in range(ncol):
            cell_p = t.cell(i, j).paragraphs[0]
            txt = r[j] if j < len(r) else ""
            crtl = is_arabic(txt, opts.lang) if txt else rtl
            if crtl:
                set_rtl_paragraph(cell_p)
            if re.fullmatch(r"[-+]?[\d.,%()\[\]\s<>=*\\]+", txt or "x"):
                cell_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if keep and i < len(cells) - 1:
                cell_p.paragraph_format.keep_with_next = True   # short tables stay on one page
            add_inline(cell_p, txt, opts, crtl, base_bold=(i == 0))


CAPTION = re.compile(r"^\**\s*(table|figure|note\.|جدول|الجدول|شكل|الشكل|ملاحظة|ملحوظة)", re.I)
TABLE_CAP = re.compile(r"^\**\s*(table|جدول|الجدول)", re.I)
FIGURE_CAP = re.compile(r"^\**\s*(figure|شكل|الشكل)", re.I)
PAPER = {"a4": (21.0, 29.7), "letter": (21.59, 27.94)}

# ---- presets, cover page, front matter ----

FM_RE = re.compile(r"\A﻿?---[ \t]*\r?\n(.*?)\r?\n---[ \t]*\r?\n", re.S)
BODY_MARK = "\x02BODY"
FRONT_HEADINGS = re.compile(r"^(abstract|الملخص|ملخص|ملخص الدراسة|dedication|الإهداء|إهداء|acknowledg(e)?ments?|"
                            r"الشكر والتقدير|شكر وتقدير|شكر و تقدير|list of abbreviations|abbreviations|قائمة الاختصارات|"
                            r"الاختصارات|declaration|authorization|التفويض|committee decision|قرار لجنة المناقشة|"
                            r"table of contents|contents|فهرس المحتويات|المحتويات|list of tables|قائمة الجداول|"
                            r"list of figures|قائمة الأشكال|list of appendices|قائمة الملاحق|preface|تمهيد)\s*:?.*$", re.I)
LABELS = {
    "en": {"by": "By", "supervisor": "Supervisor", "co_supervisor": "Co-Supervisor",
           "statement": "A thesis submitted in partial fulfillment of the requirements for the degree of {degree}",
           "toc": "Table of Contents", "tables": "List of Tables", "figures": "List of Figures",
           "fill": "Open this file in Word and accept the update prompt (or press F9 here) to build this list."},
    "ar": {"by": "إعداد", "supervisor": "إشراف", "co_supervisor": "المشرف المشارك",
           "statement": "قدمت هذه الرسالة استكمالاً لمتطلبات الحصول على درجة {degree}",
           "toc": "فهرس المحتويات", "tables": "قائمة الجداول", "figures": "قائمة الأشكال",
           "fill": "افتح الملف في Word واقبل طلب تحديث الحقول (أو اضغط F9 هنا) لبناء هذه القائمة."},
}
DEFAULT_COVER = ["logo", "university", "faculty", "title", "author", "supervisor", "co_supervisor", "statement",
                 "department", "date"]


def load_presets(path=PRESETS_FILE):
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    return {k: v for k, v in data.items() if not k.startswith("_")}


def split_front_matter(md):
    m = FM_RE.match(md)
    if not m:
        return {}, md
    meta = {}
    for line in m.group(1).splitlines():
        if ":" in line and not line.lstrip().startswith("#"):
            k, v = line.split(":", 1)
            k = k.strip().lower().replace("-", "_").replace(" ", "_")
            v = v.strip().strip('"').strip("'")
            if k and v:
                meta[k] = v
    return meta, md[m.end():]


def doc_is_arabic(meta, md, lang):
    if lang in ("ar", "en"):
        return lang == "ar"
    if meta.get("title"):
        return is_arabic(meta["title"], "auto")
    for line in md.splitlines():
        if line.startswith("#"):
            return is_arabic(line, "auto")
    return False


def cover_line(doc, text, opts, rtl, size, bold=False, caps=False, space_after=12):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = opts.cover_spacing
    if rtl:
        set_rtl_paragraph(p)
    if caps and not rtl:
        text = text.upper()
    put_runs(p, text, opts, rtl, bold=bold)
    for r in p.runs:
        r.font.size = Pt(size)
    return p


def add_cover(doc, meta, opts, rtl, md_dir):
    """Render the cover page from the front-matter block, fields in the preset's order."""
    pr = opts.preset
    lab = dict(LABELS["ar" if rtl else "en"])
    lab.update(pr.get("labels_ar" if rtl else "labels_en", {}))      # e.g. JUST says Advisor, QU says BY
    fields = pr.get("cover", DEFAULT_COVER)
    sizes = {"title": opts.size + 6, "name": opts.size + 2, "other": opts.size}
    sizes.update(pr.get("cover_sizes", {}))
    if rtl:
        sizes = {k: v + (opts.ar_size - opts.size) for k, v in sizes.items()} if not pr.get("cover_sizes") else sizes
    for _ in range(int(pr.get("cover_space_before_lines", 0))):
        cover_line(doc, "", opts, rtl, sizes["other"])
    rendered = []
    for field in fields:
        if field == "logo":
            path = meta.get("logo")
            if path:
                full = path if os.path.isabs(path) else os.path.join(md_dir, path)
                if os.path.exists(full):
                    doc.add_picture(full, width=Cm(float(pr.get("logo_width_cm", 3.0))))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    rendered.append("logo")
                else:
                    print(f"WARNING: logo not found: {full}")
            continue
        if field == "statement":
            tmpl = meta.get("statement") or pr.get("statement_ar" if rtl else "statement_en") or lab["statement"]
            # \n in the template splits lines; a line whose placeholder has no value is dropped
            # (KSU's "Major in {program}" disappears when no program is given).
            parts = []
            for part in tmpl.split("\\n"):
                missing = [k for k in ("degree", "program", "university", "faculty") if "{" + k + "}" in part and not meta.get(k)]
                if missing:
                    continue
                for k in ("degree", "program", "university", "faculty"):
                    part = part.replace("{" + k + "}", meta.get(k, ""))
                parts.append(part.strip())
            if not parts:
                continue
            for part in parts:
                cover_line(doc, part, opts, rtl, sizes["other"])
            rendered.append("statement")
            continue
        value = meta.get(field)
        if not value:
            continue
        if field == "title":
            cover_line(doc, value, opts, rtl, sizes["title"], bold=pr.get("cover_title_bold", True),
                       caps=pr.get("cover_title_caps", False), space_after=24)
        elif field in ("author", "supervisor", "co_supervisor"):
            if pr.get("cover_labels", True):
                cover_line(doc, lab["by" if field == "author" else field], opts, rtl, sizes["other"], space_after=2)
            cover_line(doc, value, opts, rtl, sizes["name"], bold=(field == "author"))
        else:
            cover_line(doc, value, opts, rtl, sizes["other"])
        rendered.append(field)
    return rendered


def add_field_block(doc, heading, instr, opts, rtl):
    """A heading paragraph plus a w:fldSimple field Word fills when fields are updated."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.keep_with_next = True
    if rtl:
        set_rtl_paragraph(p)
    put_runs(p, heading, opts, rtl, bold=True)
    for r in p.runs:
        r.font.size = Pt(opts.heading_sizes.get("h1", (opts.ar_size if rtl else opts.size) + 2))
    fp = doc.add_paragraph()
    if rtl:
        set_rtl_paragraph(fp)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = LABELS["ar" if rtl else "en"]["fill"]
    r.append(t)
    fld.append(r)
    fp._p.append(fld)
    doc.add_page_break()


def ensure_style(doc, name, rtl, opts, base="Normal"):
    try:
        st = doc.styles[name]
    except KeyError:
        st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles[base]
    if rtl:
        ppr = st.element.get_or_add_pPr()
        if ppr.find(qn("w:bidi")) is None:
            ppr.append(OxmlElement("w:bidi"))
        st.font.name = opts.ar_font
        rpr = st.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        fonts.set(qn("w:cs"), opts.ar_font)
        rpr.append(OxmlElement("w:rtl"))
    return st


def start_section(doc, opts, fmt, start, title_page=False):
    """New page + new section; page size and margins are cloned from the previous section."""
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    for old in s._sectPr.findall(qn("w:titlePg")):
        s._sectPr.remove(old)
    set_page_numbering(s, fmt, start)
    if title_page:
        s.different_first_page_header_footer = True
    return s


def build(md, opts, md_dir="."):
    doc = Document()
    pr = opts.preset
    meta, md = split_front_matter(md)
    doc_rtl = doc_is_arabic(meta, md, opts.lang)
    pw, ph = PAPER[opts.paper]
    m = dict(opts.margins)
    if pr.get("swap_for_arabic") and doc_rtl and not opts.margins_explicit:
        m["left"], m["right"] = m["right"], m["left"]
    opts.text_width_cm = pw - m["left"] - m["right"]
    pos = pr.get("page_numbers", {}).get("position", "bottom-center")
    for s in doc.sections:
        s.page_width, s.page_height = Cm(pw), Cm(ph)
        s.top_margin, s.bottom_margin = Cm(m["top"]), Cm(m["bottom"])
        s.left_margin, s.right_margin = Cm(m["left"]), Cm(m["right"])
        add_page_number(s, pos)
    if opts.mirror_margins:
        set_setting(doc, "w:mirrorMargins")
    normal = doc.styles["Normal"]
    normal.font.name = opts.font
    normal.font.size = Pt(opts.size)

    # ---- front matter: cover, TOC, lists (preset or front-matter block) ----
    md = md.replace("<!-- body -->", "\n" + BODY_MARK + "\n")
    md = re.sub(r"<!--.*?-->", "", md, flags=re.S)        # authoring notes never reach Word
    lab = LABELS["ar" if doc_rtl else "en"]
    pn = pr.get("page_numbers", {})
    fm_fmt = pn.get("front_matter", "auto" if pr else "none")
    if fm_fmt == "auto":
        fm_fmt = "abjad" if doc_rtl else "roman"
    want_toc = bool(pr.get("toc")) and not opts.no_toc
    want_lists = [x for x in pr.get("lists", []) if x in ("tables", "figures")] if not opts.no_toc else []
    cover_fields = []
    front = bool(meta.get("title")) or want_toc or bool(want_lists)
    opts.summary = {"front_matter_numbering": None, "body_section": False, "cover": [], "toc": want_toc, "lists": want_lists}
    if front:
        first = doc.sections[0]
        if fm_fmt != "none":
            set_page_numbering(first, fm_fmt, 1)
        if meta.get("title"):
            if not pn.get("title_page_number", False):
                first.different_first_page_header_footer = True     # title page counts but shows no number
            cover_fields = add_cover(doc, meta, opts, doc_rtl, md_dir)
            doc.add_page_break()
        opts.summary["cover"] = cover_fields
        opts.summary["front_matter_numbering"] = fm_fmt
        if want_toc:
            ensure_style(doc, "toc 1", doc_rtl, opts)
            ensure_style(doc, "toc 2", doc_rtl, opts)
            ensure_style(doc, "toc 3", doc_rtl, opts)
            add_field_block(doc, lab["toc"], 'TOC \\o "1-3" \\h \\z \\u', opts, doc_rtl)
        for kind in want_lists:
            cap_style = "Table Caption" if kind == "tables" else "Figure Caption"
            ensure_style(doc, cap_style, doc_rtl, opts)
            add_field_block(doc, lab[kind], 'TOC \\h \\z \\t "' + cap_style + ',1"', opts, doc_rtl)
        if want_toc or want_lists:
            set_setting(doc, "w:updateFields")        # Word offers to fill the fields when the file opens
    in_front = front and fm_fmt != "none"
    heading_shift = 1 if pr else 0                     # preset: ## -> Heading 1 (chapter), ### -> Heading 2 ...

    lines = md.splitlines()
    in_refs, i, para = False, 0, []
    after_table = False
    title_next = False          # the italic title line that follows 'Table 1' is part of the caption
    body_started = [False]

    def start_body():
        nonlocal in_front
        if in_front and not body_started[0]:
            start_section(doc, opts, "decimal", 1)
            opts.summary["body_section"] = True
        body_started[0] = True
        in_front = False

    def caption_style(p, text):
        if not want_lists or title_next:
            return
        if "tables" in want_lists and TABLE_CAP.match(text):
            p.style = doc.styles["Table Caption"]
        elif "figures" in want_lists and FIGURE_CAP.match(text):
            p.style = doc.styles["Figure Caption"]

    def flush():
        nonlocal para, after_table, title_next
        if not para:
            return
        text = " ".join(x.strip() for x in para)
        rtl = is_arabic(text, opts.lang)
        p = doc.add_paragraph()
        caption = bool(CAPTION.match(text)) or title_next
        is_label = bool(re.match(r"^\**\s*(table|figure|جدول|الجدول|شكل|الشكل)\s*\d+\**\s*$", text, re.I))
        # Table/figure labels and notes sit flush with the margin (APA), not indented.
        fmt_paragraph(p, opts, rtl, hanging=in_refs, first_line=not in_refs and opts.indent and not caption)
        if caption and not text.lower().lstrip("*").startswith(("note", "ملاحظة", "ملحوظة")):
            p.paragraph_format.keep_with_next = True             # a table title stays with its table
            caption_style(p, text)
        title_next = is_label
        if after_table:
            p.paragraph_format.space_before = Pt(4 if caption else 12)
            after_table = False
        add_inline(p, text, opts, rtl)
        para = []

    while i < len(lines):
        line = lines[i]
        if line.strip() == BODY_MARK:
            flush()
            start_body()
            i += 1
            continue
        h = re.match(r"^(#{1,6})\s+(.*)$", line)
        if h:
            flush()
            level, text = len(h.group(1)), h.group(2).strip()
            if in_front and level == 2 and not FRONT_HEADINGS.match(text):
                start_body()
            in_refs = bool(REF_HEADINGS.match(text))
            rtl = is_arabic(text, opts.lang)
            is_caption = bool(CAPTION.match(text))
            if level > 1 and not (is_caption and want_lists):
                p = doc.add_heading(level=min(max(level - heading_shift, 1), 4))
            else:
                p = doc.add_paragraph()
            if is_caption:
                caption_style(p, text)
            p.paragraph_format.keep_with_next = True
            if after_table:
                p.paragraph_format.space_before = Pt(12)
                after_table = False
            if pr and level == 2 and not is_caption and not FRONT_HEADINGS.match(text):
                if pr.get("chapter_new_page"):
                    p.paragraph_format.page_break_before = True
                if pr.get("chapter_space_before_cm"):
                    p.paragraph_format.space_before = Cm(float(pr["chapter_space_before_cm"]))
            if rtl:
                set_rtl_paragraph(p)
            centre = level == 1 or (f"h{level - heading_shift}" in pr.get("heading_center", []))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centre else WD_ALIGN_PARAGRAPH.JUSTIFY
            if pr.get("heading_caps") and not rtl and level > 1:
                text = text.upper()
            add_inline(p, text, opts, rtl, base_bold=pr.get("heading_bold", True))
            if level > 1 and not pr.get("heading_bold", True):
                for r in p.runs:                 # the Heading style is bold; KSU wants plain capitals
                    r.bold = False
            base = opts.ar_size if rtl else opts.size
            size = base + (4 if level == 1 else 2 if level == 2 else 1)
            if level > 1 and opts.heading_sizes:
                size = opts.heading_sizes.get(f"h{min(level - heading_shift, 3)}", size)
                if rtl and opts.ar_heading_sizes:
                    size = opts.ar_heading_sizes.get(f"h{min(level - heading_shift, 3)}", size)
            for r in p.runs:
                r.font.size = Pt(size)
                r.font.color.rgb = RGBColor(0, 0, 0)
            i += 1
            continue
        if line.strip().startswith("|"):
            flush()
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block, opts, is_arabic(" ".join(block), opts.lang))
            after_table = True
            continue
        li = re.match(r"^\s*([-*+]|\d+[.)])\s+(.*)$", line)
        if li:
            flush()
            text = li.group(2)
            rtl = is_arabic(text, opts.lang)
            style = "List Number" if li.group(1)[0].isdigit() else "List Bullet"
            p = doc.add_paragraph(style=style)
            p.paragraph_format.line_spacing = line_spacing(opts, rtl)
            if rtl:
                set_rtl_paragraph(p)
            add_inline(p, text, opts, rtl)
            i += 1
            continue
        if line.startswith(">"):
            flush()
            text = line.lstrip("> ").strip()
            rtl = is_arabic(text, opts.lang)
            p = doc.add_paragraph()
            fmt_paragraph(p, opts, rtl)
            p.paragraph_format.left_indent = Cm(1.27)
            p.paragraph_format.right_indent = Cm(1.27)
            add_inline(p, text, opts, rtl)
            i += 1
            continue
        if not line.strip():
            flush()
        else:
            para.append(line)
        i += 1
    flush()
    if in_front and front:
        print("WARNING: no body start found (no <!-- body --> and no ## chapter heading): "
              "the whole file carries front-matter page numbers.")
    return doc


def resolve_options(opts, pr):
    """CLI flag beats preset beats built-in default."""
    defaults = {"font": "Times New Roman", "ar_font": "Simplified Arabic", "size": 12.0, "ar_size": 14.0,
                "spacing": 2.0, "paper": "a4", "table_style": "apa"}
    for k, d in defaults.items():
        if getattr(opts, k) is None:
            setattr(opts, k, pr.get(k, d))
    opts.ar_spacing = pr.get("ar_spacing")
    opts.refs_spacing = pr.get("refs_spacing")
    opts.justify = pr.get("justify", True)
    opts.heading_sizes = pr.get("heading_sizes", {})
    opts.ar_heading_sizes = pr.get("ar_heading_sizes", {})
    opts.cover_spacing = pr.get("cover_spacing", 1.5)
    if pr.get("indent") is False and opts.indent:
        opts.indent = False
    opts.margins_explicit = bool(opts.margin_cm or opts.margins_cm)
    if opts.margins_cm:
        t, b, l, r = [float(x) for x in opts.margins_cm.split(",")]
        opts.margins = {"top": t, "bottom": b, "left": l, "right": r}
    elif opts.margin_cm:
        opts.margins = {k: opts.margin_cm for k in ("top", "bottom", "left", "right")}
    elif pr.get("margins_cm"):
        opts.margins = {k: float(pr["margins_cm"][k]) for k in ("top", "bottom", "left", "right")}
    else:
        opts.margins = {k: 2.54 for k in ("top", "bottom", "left", "right")}
    if opts.mirror_margins is None:
        opts.mirror_margins = bool(pr.get("mirror_margins", False))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("markdown", nargs="?")
    ap.add_argument("-o", "--out")
    ap.add_argument("--lang", choices=["auto", "ar", "en"], default="auto",
                    help="auto decides per paragraph from its script")
    ap.add_argument("--font", default=None, help="Latin font (default Times New Roman)")
    ap.add_argument("--ar-font", default=None, help="Arabic font (default Simplified Arabic)")
    ap.add_argument("--size", type=float, default=None, help="Latin size pt (default 12)")
    ap.add_argument("--ar-size", type=float, default=None, help="Arabic size pt (default 14)")
    ap.add_argument("--spacing", type=float, default=None, help="APA uses 2.0 (default); many Arab universities ask for 1.5")
    ap.add_argument("--margin-cm", type=float, default=None, help="all four margins (default 2.54)")
    ap.add_argument("--margins-cm", default=None, metavar="T,B,L,R", help="asymmetric margins, e.g. 2.5,2.5,3.5,2.5 (binding left)")
    ap.add_argument("--mirror-margins", action="store_true", default=None,
                    help="Word mirror margins: left/right become inside/outside for double-sided printing")
    ap.add_argument("--no-indent", dest="indent", action="store_false", help="no first-line indent on body paragraphs")
    ap.add_argument("--table-style", choices=["apa", "grid"], default=None,
                    help="apa = horizontal rules only (APA 7, default); grid = all borders (some universities)")
    ap.add_argument("--paper", choices=["a4", "letter"], default=None,
                    help="a4 (Arab and most non-US universities, default) or letter (US journals)")
    ap.add_argument("--preset", default=None, help="university preset name from references/university-presets.json")
    ap.add_argument("--no-toc", action="store_true", help="skip the preset's TOC and lists")
    ap.add_argument("--list-presets", action="store_true", help="show the available presets and exit")
    opts = ap.parse_args()

    presets = {}
    if opts.preset or opts.list_presets:
        try:
            presets = load_presets()
        except (OSError, ValueError) as e:
            sys.exit(f"cannot read presets file {PRESETS_FILE}: {e}")
    if opts.list_presets:
        print("| preset | university | verified | paper | binding margin | fonts | spacing |")
        print("|---|---|---|---|---|---|---|")
        for name, p in presets.items():
            mg = p.get("margins_cm", {})
            v = ("yes " + p.get("verified_date", "")) if p.get("verified") else "no (common practice)"
            print(f"| {name} | {p.get('label', '')} | {v} | {p.get('paper', 'a4')} | "
                  f"L {mg.get('left', 2.54)} / R {mg.get('right', 2.54)} cm | {p.get('ar_font', 'Simplified Arabic')} "
                  f"{p.get('ar_size', 14)} / {p.get('font', 'Times New Roman')} {p.get('size', 12)} | {p.get('spacing', 2.0)} |")
        return
    if not opts.markdown or not opts.out:
        ap.error("markdown and -o/--out are required")
    pr = {}
    if opts.preset:
        if opts.preset not in presets:
            sys.exit(f"unknown preset '{opts.preset}'. Available: {', '.join(presets)}")
        pr = presets[opts.preset]
    opts.preset = pr
    resolve_options(opts, pr)
    with open(opts.markdown, encoding="utf-8") as f:
        md = f.read()
    doc = build(md, opts, os.path.dirname(os.path.abspath(opts.markdown)))
    doc.save(opts.out)
    print(f"saved {opts.out}")
    if pr:
        v = f"verified against the university guide on {pr.get('verified_date', '?')}" if pr.get("verified") \
            else "UNVERIFIED: based on common practice, confirm with the department"
        print(f"preset: {pr.get('label', '')} | {v}")
        print(f"source: {pr.get('source_url', 'none')}")
        if pr.get("notes"):
            print(f"notes: {pr['notes']}")
    s = getattr(opts, "summary", {})
    mg = opts.margins
    print(f"margins cm: top {mg['top']} bottom {mg['bottom']} left {mg['left']} right {mg['right']}"
          f"{' (swapped: Arabic thesis, binding on the right)' if pr.get('swap_for_arabic') and mg['right'] > mg['left'] else ''}"
          f" | mirror margins: {'on' if opts.mirror_margins else 'off'}")
    if s.get("cover"):
        print(f"cover page: {', '.join(s['cover'])}")
    if s.get("toc") or s.get("lists"):
        print("TOC / lists: fields inserted. Word fills them: open the file and accept the update prompt, "
              "or Ctrl+A then F9. (With Word COM: $d.Fields.Update() before SaveAs.)")
    if s.get("front_matter_numbering") and s["front_matter_numbering"] != "none":
        body = "body restarts at 1 (section break inserted automatically)" if s.get("body_section") \
            else "no body section break was inserted"
        print(f"page numbers: front matter {NUM_FMT.get(s['front_matter_numbering'], s['front_matter_numbering'])}, {body}")


if __name__ == "__main__":
    main()
