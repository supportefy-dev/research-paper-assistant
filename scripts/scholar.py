#!/usr/bin/env python3
"""scholar.py - verified-source toolkit for the research-paper-assistant skill.

Every source that reaches a manuscript passes through this tool, so nothing is
cited from memory. All APIs used are free.

Commands
  search    find candidate papers (OpenAlex, Semantic Scholar, Crossref for Arabic journals); --issn for scope checks
  verify    check DOIs and retractions, match one citation, or a whole list (--citations-file)
  add       verify a DOI or ISBN and append it to the source ledger
  fix       correct a ledger entry's author split, title case or year (never its identity)
  book      look up a book by ISBN or title (Open Library)
  abstract  print the real abstract of a DOI
  cite      formatted reference for a DOI
  refs      build the reference list from the ledger (only what the text cites); --format ris/bibtex/csl for Zotero
  list      show the ledger
  audit     check a manuscript's citations and reference list against the ledger
  numbers   check every statistic in a text against labelled analysis output
  journal   vet one journal (SJR quartiles, OpenAlex stats, DOAJ, hijack list)
  journals  shortlist journals from an SJR CSV by category and quartile

Optional environment variables
  RESEARCH_MAILTO   contact email, puts Crossref calls in the faster polite pool
  OPENALEX_API_KEY  free key from openalex.org/settings/api ($1/day of credits)
  S2_API_KEY        free Semantic Scholar key (request form on their API page)
  SJR_CSV           path to a SCImago CSV the user downloaded in a browser

Run with:  python -X utf8 scholar.py <command> --help
"""

import argparse
import csv
import datetime as dt
import difflib
import html
import json
import os
import re
import sys
import time
import unicodedata
import urllib.parse

try:
    import requests
except ImportError:
    sys.exit("requests is missing: pip install requests")

for stream in (sys.stdout, sys.stderr):
    try:
        stream.reconfigure(encoding="utf-8")
    except (AttributeError, ValueError):
        pass

MAILTO = os.environ.get("RESEARCH_MAILTO", "").strip()
# Private key files live outside every repo: ~/.config/research-paper-assistant/<name>.key
KEY_DIR = os.path.join(os.path.expanduser("~"), ".config", "research-paper-assistant")


def _secret(env, name):
    val = os.environ.get(env, "").strip()
    path = os.path.join(KEY_DIR, name)
    if not val and os.path.exists(path):
        with open(path, encoding="utf-8") as f:
            val = f.read().strip()
    return val


OPENALEX_KEY = _secret("OPENALEX_API_KEY", "openalex.key")
S2_KEY = os.environ.get("S2_API_KEY", "").strip()
TODAY = dt.date.today().isoformat()
UA = "research-paper-assistant/1.0" + (f" (mailto:{MAILTO})" if MAILTO else "")

INTEGRITY_FLAGS = {"retraction", "partial_retraction", "withdrawal", "removal", "expression_of_concern"}
SOFT_FLAGS = {"correction", "erratum", "corrigendum", "addendum"}
# Hyphen variants, en/em dashes, horizontal bar, minus sign: all become "-".
DASHES = dict.fromkeys([0x2010, 0x2011, 0x2012, 0x2013, 0x2014, 0x2015, 0x2212], "-")
EN_EM = chr(0x2013) + chr(0x2014)
ARABIC_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹٫", "01234567890123456789.")
ARABIC_RE = re.compile(r"[؀-ۿ]")

# ---------------------------------------------------------------- http

_last_call = {}
# Minimum seconds between calls per host. Crossref list endpoints allow 1/s
# (public) or 3/s (polite); Semantic Scholar's shared pool throttles hard.
_INTERVAL = {
    "api.crossref.org": 0.35 if MAILTO else 1.05,
    "api.semanticscholar.org": 1.1,
    "doaj.org": 0.55,
    "api.openalex.org": 0.12,
    "openlibrary.org": 0.4,
    "doi.org": 0.35,
}


class HTTPError(Exception):
    def __init__(self, status, url, body=""):
        super().__init__(f"HTTP {status} for {url}")
        self.status, self.url, self.body = status, url, body


def http_get(url, params=None, headers=None, retries=4, timeout=40, allow_404=True):
    """GET with per-host pacing and exponential backoff on 429/5xx.

    Returns the Response, or None for a 404 when allow_404 is set. None of the
    APIs send Retry-After on 429, so the client paces itself.
    """
    host = urllib.parse.urlparse(url).netloc
    hdrs = {"User-Agent": UA}
    hdrs.update(headers or {})
    delay = 2.0
    for attempt in range(retries + 1):
        gap = _INTERVAL.get(host, 0.2) - (time.time() - _last_call.get(host, 0))
        if gap > 0:
            time.sleep(gap)
        _last_call[host] = time.time()
        try:
            r = requests.get(url, params=params, headers=hdrs, timeout=timeout)
        except requests.RequestException as e:
            if attempt == retries:
                raise HTTPError("network", url, str(e))
            time.sleep(delay)
            delay *= 2
            continue
        if r.status_code == 404 and allow_404:
            return None
        if r.status_code in (429, 500, 502, 503, 504) and attempt < retries:
            time.sleep(delay)
            delay *= 2
            continue
        if r.status_code >= 400:
            raise HTTPError(r.status_code, url, r.text[:300])
        return r
    raise HTTPError("retries", url)


def as_json(r):
    return json.loads(r.content.decode("utf-8"))


# ---------------------------------------------------------------- text helpers

INVISIBLE = re.compile("[\u200b-\u200f\u202a-\u202e\u2066-\u2069\ufeff]")
QUOTES = dict.fromkeys([0x2018, 0x2019, 0x201B, 0x2032], "'") | dict.fromkeys([0x201C, 0x201D, 0x2033], '"')


def clean(s):
    """Unescape HTML, drop tags and invisible bidi marks, ASCII hyphens, one space."""
    if s is None:
        return ""
    s = html.unescape(str(s))
    s = re.sub(r"<[^>]+>", "", s)
    s = INVISIBLE.sub("", s).translate(DASHES)
    return re.sub(r"\s+", " ", s).strip()


def cmp_text(s):
    """Comparison form: quotes straightened, no whitespace, case-folded."""
    return re.sub(r"\s+", "", clean(s).translate(QUOTES)).casefold()


def norm_doi(s):
    s = (s or "").strip().strip(".,;")
    s = re.sub(r"^(https?://(dx\.)?doi\.org/|doi:\s*)", "", s, flags=re.I)
    return s


def name_key(s):
    """Comparable form of a surname: no accents, no al-/el- prefix, letters only."""
    s = unicodedata.normalize("NFKD", clean(s)).lower()
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"^(al|el)[\s\-']+", "", s)
    s = re.sub(r"^ال", "", s)
    return re.sub(r"[^a-z؀-ۿ]", "", s)


def title_sim(a, b):
    a, b = clean(a).lower(), clean(b).lower()
    return difflib.SequenceMatcher(None, a, b).ratio()


def compact_issn(s):
    s = re.sub(r"[^0-9Xx]", "", s or "").upper()
    return s if len(s) == 8 else ""


def hyphen_issn(s):
    c = compact_issn(s)
    return f"{c[:4]}-{c[4:]}" if c else ""


def short(s, n):
    s = clean(s)
    return s if len(s) <= n else s[: n - 3] + "..."


def table(rows, headers):
    """Plain pipe table, readable in a terminal and pasteable as markdown."""
    rows = [[clean(c) for c in r] for r in rows]
    out = ["| " + " | ".join(headers) + " |", "|" + "|".join("---" for _ in headers) + "|"]
    out += ["| " + " | ".join(r) + " |" for r in rows]
    return "\n".join(out)


# ---------------------------------------------------------------- OpenAlex

_OA_STATE = {"key_ok": True}


def openalex(path, params=None):
    params = dict(params or {})
    use_key = OPENALEX_KEY and _OA_STATE["key_ok"]
    if use_key:
        params["api_key"] = OPENALEX_KEY
    try:
        r = http_get("https://api.openalex.org/" + path, params=params)
    except HTTPError as e:
        if use_key and e.status in (401, 403):
            # Not activated yet, revoked or mistyped: say so once and carry on keyless.
            _OA_STATE["key_ok"] = False
            print(f"[openalex] the API key was refused (HTTP {e.status}: not activated yet, or wrong). "
                  f"Continuing without it. Key file: {os.path.join(KEY_DIR, 'openalex.key')}", file=sys.stderr)
            params.pop("api_key", None)
            r = http_get("https://api.openalex.org/" + path, params=params)
        else:
            raise
    if r is None:
        return None
    left = r.headers.get("X-RateLimit-Remaining")
    if left is not None and int(float(left)) < 150:
        print(f"[openalex] only {left} credits left today (a search costs 10). "
              "A free key at openalex.org/settings/api gives 10,000/day.", file=sys.stderr)
    return as_json(r)


def rebuild_abstract(inv):
    if not inv:
        return ""
    pos = {i: w for w, idx in inv.items() for i in idx}
    return " ".join(pos[i] for i in sorted(pos))


def openalex_by_doi(doi):
    """Single-entity lookup: costs 0 credits."""
    return openalex("works/doi:" + urllib.parse.quote(doi, safe="/()"))


def oa_row(w):
    loc = w.get("primary_location") or {}
    src = loc.get("source") or {}
    auths = w.get("authorships") or []
    first = ((auths[0].get("author") or {}).get("display_name", "")) if auths else ""
    return {
        "doi": norm_doi(w.get("doi")),
        "title": clean(w.get("display_name") or w.get("title")),
        "year": w.get("publication_year"),
        "cites": w.get("cited_by_count"),
        "first_author": first,
        "authors": [(x.get("author") or {}).get("display_name", "") for x in auths],
        "n_authors": len(auths),
        "venue": clean(src.get("display_name")),
        "issn_l": src.get("issn_l"),
        "core_venue": src.get("is_core"),
        "type": w.get("type"),
        "is_oa": (w.get("open_access") or {}).get("is_oa"),
        "retracted": w.get("is_retracted"),
        "abstract": rebuild_abstract(w.get("abstract_inverted_index")),
        "source": "openalex",
    }


def search_openalex(q, a):
    filters = ["is_retracted:false"]
    if q:
        filters.append(f"title_and_abstract.search:{q}")
    if a.issn:
        filters.append("primary_location.source.issn:" + "|".join(hyphen_issn(x) for x in a.issn.split(",")))
    if a.core:
        filters.append("primary_location.source.is_core:true")
    if a.has_doi:
        filters.append("has_doi:true")
    if a.since:
        filters.append(f"from_publication_date:{a.since}-01-01")
    if a.until:
        filters.append(f"to_publication_date:{a.until}-12-31")
    if a.type:
        filters.append(f"type:{a.type}")
    if a.min_cites:
        filters.append(f"cited_by_count:>{a.min_cites - 1}")
    if a.oa:
        filters.append("is_oa:true")
    if a.lang:
        filters.append(f"language:{a.lang}")
    params = {
        "filter": ",".join(filters),
        "per_page": min(a.limit, 200),
        "select": "id,doi,display_name,publication_year,cited_by_count,primary_location,"
                  "authorships,type,open_access,is_retracted"
                  + (",abstract_inverted_index" if a.abstracts else ""),
    }
    if a.sort == "cites":
        params["sort"] = "cited_by_count:desc"
    elif a.sort == "recent":
        params["sort"] = "publication_date:desc"
    data = openalex("works", params)
    return [oa_row(w) for w in (data or {}).get("results", [])], (data or {}).get("meta", {}).get("count")


# ---------------------------------------------------------------- Semantic Scholar

S2_FIELDS = "title,year,authors,externalIds,citationCount,abstract,venue,publicationTypes,isOpenAccess"


def s2(path, params=None):
    headers = {"x-api-key": S2_KEY} if S2_KEY else {}
    r = http_get("https://api.semanticscholar.org/graph/v1/" + path, params=params,
                 headers=headers, retries=6)
    return as_json(r) if r is not None else None


def s2_row(p):
    ext = p.get("externalIds") or {}
    auths = p.get("authors") or []
    return {
        "doi": norm_doi(ext.get("DOI")),
        "title": clean(p.get("title")),
        "year": p.get("year"),
        "cites": p.get("citationCount"),
        "first_author": auths[0]["name"] if auths else "",
        "authors": [x.get("name", "") for x in auths],
        "n_authors": len(auths),
        "venue": clean(p.get("venue")),
        "issn_l": None,
        "core_venue": None,
        "type": ",".join(p.get("publicationTypes") or []),
        "is_oa": p.get("isOpenAccess"),
        "retracted": None,
        "abstract": clean(p.get("abstract")),
        "source": "semanticscholar",
    }


def search_s2(q, a):
    params = {"query": q, "limit": min(a.limit, 100), "fields": S2_FIELDS}
    if a.since or a.until:
        params["year"] = f"{a.since or ''}-{a.until or ''}"
    if a.min_cites:
        params["minCitationCount"] = a.min_cites
    if a.type == "article":
        params["publicationTypes"] = "JournalArticle"
    if a.oa:
        params["openAccessPdf"] = ""
    data = s2("paper/search", params) or {}
    return [s2_row(p) for p in data.get("data", [])], data.get("total")


S2_STATE = {"throttled": False}


def s2_by_doi(doi):
    try:
        return s2("paper/DOI:" + doi, {"fields": S2_FIELDS})
    except HTTPError as e:
        if e.status in (429, "retries", 500, 503):
            S2_STATE["throttled"] = True
        return None


# ---------------------------------------------------------------- Crossref / DOI

def crossref_work(doi):
    params = {"mailto": MAILTO} if MAILTO else None
    r = http_get("https://api.crossref.org/works/" + urllib.parse.quote(doi, safe=""), params=params)
    return as_json(r)["message"] if r is not None else None


def handle_exists(doi):
    """doi.org Handle API: responseCode 1 = registered, 100 = not registered."""
    r = http_get("https://doi.org/api/handles/" + urllib.parse.quote(doi, safe="/"))
    if r is None:
        return False
    return as_json(r).get("responseCode") == 1


def csl_json(doi):
    r = http_get("https://doi.org/" + urllib.parse.quote(doi, safe="/()"),
                 headers={"Accept": "application/vnd.citationstyles.csl+json"})
    return as_json(r) if r is not None else None


def retraction_status(msg):
    """Read Crossref's updated-by (includes Retraction Watch data). Never use select=."""
    seen, notices = set(), []
    for u in msg.get("updated-by") or []:
        t = (u.get("type") or "").lower().replace("-", "_")
        key = (u.get("DOI", "").lower(), t)
        if key in seen:
            continue
        seen.add(key)
        parts = ((u.get("updated") or {}).get("date-parts") or [[None]])[0]
        notices.append({"type": t, "notice_doi": u.get("DOI"), "source": u.get("source"),
                        "year": parts[0] if parts else None})
    types = {n["type"] for n in notices}
    if types & {"retraction", "partial_retraction", "withdrawal", "removal"}:
        status = "retracted"
    elif "expression_of_concern" in types:
        status = "expression_of_concern"
    elif types & SOFT_FLAGS:
        status = "corrected"
    else:
        status = "ok"
    titles = " ".join(msg.get("title") or [])
    if status == "ok" and re.match(r"\s*(retracted|withdrawn)\b", titles, re.I):
        status = "retracted"
    is_notice = bool(msg.get("update-to"))
    return {"status": status, "notices": notices, "is_update_notice": is_notice}


def _year(d):
    parts = ((d or {}).get("date-parts") or [[None]])[0]
    return parts[0] if parts and parts[0] else None


def crossref_years(msg):
    keys = ("issued", "published", "published-print", "published-online", "created")
    ys = {_year(msg.get(k)) for k in keys} | {_year((msg.get("journal-issue") or {}).get("published-print"))}
    return sorted(y for y in ys if y)


def date_year(msg):
    """APA cites the year of the final version. 'issued' is the earliest date
    (online-first), so once a paper has a volume, take the print/issue year."""
    if msg.get("volume"):
        y = _year((msg.get("journal-issue") or {}).get("published-print")) or _year(msg.get("published-print"))
        if y:
            return y
    for k in ("issued", "published", "published-print", "published-online", "created"):
        y = _year(msg.get(k))
        if y:
            return y
    return None


def first(v):
    if isinstance(v, list):
        return v[0] if v else ""
    return v or ""


def record_from_crossref(msg):
    authors = []
    marks = r"[\s*\u2020\u2021§¹²³⁴]+$"            # affiliation markers some publishers leave in names
    for a in msg.get("author") or []:
        if a.get("family"):
            authors.append({"family": re.sub(marks, "", clean(a["family"])),
                            "given": re.sub(marks, "", clean(a.get("given", "")))})
        elif a.get("name"):
            authors.append({"family": clean(a["name"]), "given": "", "org": True})
    return {
        "type": msg.get("type"),
        "doi": msg.get("DOI"),
        "title": clean(first(msg.get("title"))),
        "authors": authors,
        "year": date_year(msg),
        "container": clean(first(msg.get("container-title"))),
        "volume": clean(msg.get("volume")),
        "issue": clean(msg.get("issue")),
        "pages": clean(msg.get("page")),
        "article_number": clean(msg.get("article-number")),
        "issn": msg.get("ISSN") or [],
        "publisher": clean(msg.get("publisher")),
        "years_seen": crossref_years(msg),
        "crossref_abstract": clean(msg.get("abstract")),
        "retraction": retraction_status(msg),
        "verified": {"by": "crossref", "at": TODAY},
    }


def record_from_csl(c, doi, agency):
    authors = [{"family": clean(a.get("family") or a.get("literal", "")), "given": clean(a.get("given", ""))}
               for a in c.get("author") or []]
    parts = ((c.get("issued") or {}).get("date-parts") or [[None]])[0]
    return {
        "type": c.get("type"),
        "doi": doi,
        "title": clean(first(c.get("title"))),
        "authors": authors,
        "year": parts[0] if parts else None,
        "container": clean(first(c.get("container-title"))),
        "volume": clean(c.get("volume")),
        "issue": clean(c.get("issue")),
        "pages": clean(c.get("page")),
        "article_number": "",
        "issn": [],
        "publisher": clean(c.get("publisher")),
        "retraction": {"status": "unknown", "notices": [], "is_update_notice": False},
        "verified": {"by": agency, "at": TODAY},
    }


def resolve_doi(doi):
    """Return (record, verdict). verdict: found | found-non-crossref | not-found."""
    doi = norm_doi(doi)
    msg = crossref_work(doi)
    if msg is not None:
        return record_from_crossref(msg), "found"
    # A Crossref 404 is not proof of fabrication: DataCite (datasets, arXiv,
    # Zenodo) DOIs live elsewhere. Ask the DOI registry itself.
    if handle_exists(doi):
        c = csl_json(doi)
        if c:
            return record_from_csl(c, doi, "doi.org"), "found-non-crossref"
        return None, "registered-no-metadata"
    return None, "not-found"


def crossref_query(text, rows=5):
    params = {"query.bibliographic": text, "rows": rows}
    if MAILTO:
        params["mailto"] = MAILTO
    r = http_get("https://api.crossref.org/works", params=params)
    return as_json(r)["message"]["items"] if r is not None else []


CROSSREF_SELECT = ("DOI,title,author,issued,published-print,published-online,created,container-title,abstract,"
                   "type,publisher,ISSN,volume,issue,page,article-number,is-referenced-by-count,updated-by,update-to")
# scholar --type values that Crossref can filter on (it has no 'review' type and no language filter).
CROSSREF_TYPES = {"article": "journal-article", "book-chapter": "book-chapter", "book": "book",
                  "dissertation": "dissertation", "proceedings-article": "proceedings-article", "report": "report"}


def cr_row(it):
    """Crossref work -> the same candidate row shape as oa_row()."""
    rec = record_from_crossref(it)
    auths = [a["family"] if a.get("org") else f"{a.get('given', '')} {a['family']}".strip() for a in rec["authors"]]
    issn = rec["issn"]
    return {
        "doi": norm_doi(rec["doi"]),
        "title": rec["title"],
        "year": rec["year"],
        "cites": it.get("is-referenced-by-count"),
        "first_author": auths[0] if auths else "",
        "authors": auths,
        "n_authors": len(auths),
        "venue": rec["container"],
        "issn_l": hyphen_issn(issn[0]) if issn else None,
        "core_venue": None,
        "type": rec["type"],
        "is_oa": None,
        "retracted": rec["retraction"]["status"] == "retracted",
        "abstract": re.sub(r"^(abstract|summary|الملخص|ملخص)\s*[:.]?\s*", "", rec["crossref_abstract"], flags=re.I),
        "source": "crossref",
    }


def search_crossref(q, a):
    """Crossref bibliographic search: the only free index that covers the Arabic
    journals OpenAlex misses (Egyptian Knowledge Bank 10.21608, Al Manhal 10.12816,
    Jordanian and Gulf university presses). Relevance ranking is Crossref's own and
    'total-results' counts any work sharing a token, so only the top rows matter."""
    params = {"query.bibliographic": q, "rows": min(a.limit, 100), "select": CROSSREF_SELECT}
    if MAILTO:
        params["mailto"] = MAILTO
    filters = []
    if a.type:
        if a.type in CROSSREF_TYPES:
            filters.append("type:" + CROSSREF_TYPES[a.type])
        else:
            print(f"[crossref] no type filter for '{a.type}': searching all types", file=sys.stderr)
    if a.since:
        filters.append(f"from-pub-date:{a.since}")
    if a.until:
        filters.append(f"until-pub-date:{a.until}")
    if a.issn:
        filters += ["issn:" + hyphen_issn(x) for x in a.issn.split(",") if hyphen_issn(x)]
    for flag, name in ((a.core, "--core"), (a.oa, "--oa"), (a.lang, "--lang")):
        if flag:
            print(f"[crossref] {name} is not supported by Crossref and was ignored", file=sys.stderr)
    if filters:
        params["filter"] = ",".join(filters)
    if a.sort == "cites":
        params.update(sort="is-referenced-by-count", order="desc")
    elif a.sort == "recent":
        params.update(sort="published", order="desc")
    r = http_get("https://api.crossref.org/works", params=params)
    msg = as_json(r)["message"] if r is not None else {}
    rows = [cr_row(it) for it in msg.get("items", [])]
    dropped = [r_ for r_ in rows if r_["retracted"]]
    if dropped:
        print(f"[crossref] dropped {len(dropped)} retracted work(s): " + ", ".join(x["doi"] for x in dropped), file=sys.stderr)
    rows = [r_ for r_ in rows if not r_["retracted"]]
    if a.min_cites:
        rows = [r_ for r_ in rows if (r_["cites"] or 0) >= a.min_cites]
    return rows, msg.get("total-results")


# Small, visible Arabic synonym map for --expand. A term maps to the other wordings
# Arab journals use for the SAME construct; a term that maps to nothing is kept as is
# because it is already the standard wording. Print it with `search --show-expansions`.
ARABIC_SYNONYMS = {
    "القيادة التحويلية": [],
    "الأداء الوظيفي": ["أداء العاملين", "أداء الموظفين"],
    "الرضا الوظيفي": ["رضا الموظفين", "رضا العاملين"],
    "الالتزام التنظيمي": ["الولاء التنظيمي"],
    "الثقافة التنظيمية": ["ثقافة المنظمة"],
    "التمكين الإداري": ["تمكين العاملين"],
    "الاحتراق الوظيفي": ["الاحتراق النفسي", "الإنهاك الوظيفي"],
    "التحصيل الدراسي": ["التحصيل الأكاديمي"],
    "التعليم الإلكتروني": ["التعلم الإلكتروني", "التعليم عن بعد"],
    "جودة الخدمة": ["جودة الخدمات"],
    "دوران العمل": ["ترك العمل", "نية ترك العمل"],
}
MAX_EXPANSIONS = 4


def expand_query(q):
    """Return [(term, variant, new_query)] for the mapped terms found in q, one
    substitution per variant (no cross product), capped at MAX_EXPANSIONS."""
    out, kept = [], []
    for term, variants in ARABIC_SYNONYMS.items():
        if term not in q:
            continue
        if not variants:
            kept.append(term)
        for v in variants:
            out.append((term, v, q.replace(term, v)))
    return out[:MAX_EXPANSIONS], kept


def parse_citation(text):
    """Pull year / title / journal / volume / issue / first page out of an APA-ish reference."""
    t = clean(text)
    out = {}
    m = re.search(r"\((\d{4})[a-z]?\)", t) or re.search(r"\b((?:19|20)\d{2})\b", t)
    if m:
        out["year"] = int(m.group(1))
        rest = t[m.end():].lstrip(". ")
        tm = re.match(r"(.+?[.?!])\s+(.*)$", rest)
        if tm:
            out["title"], tail = tm.group(1).rstrip("."), tm.group(2)
            jm = re.match(r"(.+?),\s*(\d+)\s*(?:\(([^)]+)\))?\s*,\s*(?:Article\s+)?([A-Za-z]?\d+)", tail)
            if jm:
                out.update(journal=jm.group(1).strip(), volume=jm.group(2), issue=(jm.group(3) or "").strip(),
                           first_page=jm.group(4))
    return out


def discrepancies(cited, rec):
    d = []
    if cited.get("year") and rec.get("year") and cited["year"] not in (rec.get("years_seen") or [rec["year"]]):
        d.append(f"year {cited['year']} -> {rec['year']}")
    if cited.get("journal") and rec.get("container") and title_sim(cited["journal"].replace("&", "and"),
                                                                  rec["container"].replace("&", "and")) < 0.75:
        d.append(f"journal '{cited['journal']}' -> '{rec['container']}'")
    if cited.get("volume") and rec.get("volume") and cited["volume"] != rec["volume"]:
        d.append(f"volume {cited['volume']} -> {rec['volume']}")
    if cited.get("issue") and rec.get("issue") and cited["issue"] != rec["issue"]:
        d.append(f"issue {cited['issue']} -> {rec['issue']}")
    fp = (rec.get("pages") or "").split("-")[0] or rec.get("article_number")
    if cited.get("first_page") and fp and cited["first_page"] not in (fp, rec.get("article_number")):
        d.append(f"pages {cited['first_page']} -> {rec.get('pages') or rec.get('article_number')}")
    return d


def match_citation(text, rows=5):
    items = crossref_query(text, rows)
    cited = parse_citation(text)
    years = re.findall(r"\b(19\d{2}|20\d{2})\b", text)
    author_seg = text[: text.find("(")] if "(" in text else text[:120]
    author_tokens = {name_key(w) for w in re.findall(r"[\w'\-؀-ۿ]+", author_seg)}
    out = []
    for it in items:
        rec = record_from_crossref(it)
        fam = rec["authors"][0]["family"] if rec["authors"] else ""
        # Crossref always returns hits, even for invented citations, so the
        # verdict rests on title similarity plus author agreement, then the
        # remaining fields are compared to report what the citation got wrong.
        t_sim = title_sim(rec["title"], cited["title"]) if cited.get("title") else max(
            (title_sim(rec["title"], seg) for seg in re.split(r"[.?!]\s", text) if len(seg) > 15), default=0.0)
        fk = name_key(fam)
        author_ok = bool(fk) and (fk in author_tokens or (len(fk) >= 5 and fk in name_key(author_seg)))
        year_ok = (rec["year"] in [int(y) for y in years] or any(int(y) in rec.get("years_seen", []) for y in years)) \
            if years else None
        diffs = discrepancies(cited, rec)
        if t_sim >= 0.85 and author_ok:
            verdict = "match" if not diffs else "match-with-errors"
        else:
            verdict = "possible" if t_sim >= 0.6 else "no-match"
        out.append({"doi": rec["doi"], "title": rec["title"], "first_author": fam, "year": rec["year"],
                    "venue": rec["container"], "title_similarity": round(t_sim, 2),
                    "author_match": author_ok, "year_match": year_ok, "verdict": verdict,
                    "discrepancies": diffs, "retraction": rec["retraction"]["status"], "record": rec})
    out.sort(key=lambda x: ({"match": 0, "match-with-errors": 1, "possible": 2, "no-match": 3}[x["verdict"]],
                            -x["title_similarity"]))
    return out, cited


def locator_owner(cited):
    """Fabricated references often borrow a real journal + volume + page.
    Return the paper that actually sits at that locator, if any."""
    if not (cited.get("journal") and cited.get("volume") and cited.get("first_page")):
        return None
    for it in crossref_query(f"{cited['journal']} {cited['volume']} {cited['first_page']}", rows=5):
        rec = record_from_crossref(it)
        fp = (rec.get("pages") or "").split("-")[0]
        if rec.get("volume") == cited["volume"] and cited["first_page"] in (fp, rec.get("article_number")) \
                and title_sim(cited["journal"].replace("&", "and"), rec["container"].replace("&", "and")) >= 0.75:
            return rec
    return None


def judge_citation(text):
    """One reference -> (verdict, best candidate or None, detail)."""
    res, cited = match_citation(text)
    best = res[0] if res else None
    if best and best["verdict"] in ("match", "match-with-errors"):
        st = best["retraction"]
        if st in ("retracted", "expression_of_concern"):
            return "RETRACTED" if st == "retracted" else "CONCERN", best, st
        if best["verdict"] == "match":
            return "REAL", best, ""
        return "REAL, ERRORS", best, "; ".join(best["discrepancies"])
    owner = locator_owner(cited)
    if owner:
        return "FABRICATED?", None, (f"its journal/volume/page belongs to a different paper: "
                                     f"'{short(owner['title'], 70)}' https://doi.org/{owner['doi']}")
    if best and best["verdict"] == "possible":
        return "UNVERIFIED", best, f"closest: '{short(best['title'], 60)}' (title similarity {best['title_similarity']})"
    return "UNVERIFIED", None, "no matching record in Crossref"


# ---------------------------------------------------------------- Open Library (books)

def book_lookup(isbn=None, title=None, author=None, limit=5):
    params = {"fields": "title,subtitle,author_name,first_publish_year,publish_year,publisher,isbn,key",
              "limit": limit}
    if isbn:
        params["isbn"] = re.sub(r"[^0-9Xx]", "", isbn)
    else:
        params["title"] = title
        if author:
            params["author"] = author
    r = http_get("https://openlibrary.org/search.json", params=params)
    docs = as_json(r).get("docs", []) if r is not None else []
    out = []
    for d in docs:
        names = d.get("author_name") or []
        authors = []
        for n in names:
            bits = clean(n).split()
            authors.append({"family": bits[-1] if bits else n, "given": " ".join(bits[:-1])})
        out.append({
            "type": "book",
            "doi": None,
            "isbn": (d.get("isbn") or [None])[0] if not isbn else re.sub(r"[^0-9Xx]", "", isbn),
            "title": clean(d.get("title") + (": " + d["subtitle"] if d.get("subtitle") else "")),
            "authors": authors,
            "year": d.get("first_publish_year"),
            "publish_years": sorted(set(d.get("publish_year") or []))[:12],
            "container": "",
            "publisher": clean((d.get("publisher") or [""])[0]),
            "publishers": [clean(p) for p in (d.get("publisher") or [])[:6]],
            "openlibrary": "https://openlibrary.org" + d.get("key", ""),
            "retraction": {"status": "n/a", "notices": [], "is_update_notice": False},
            "verified": {"by": "openlibrary", "at": TODAY},
        })
    return out


# ---------------------------------------------------------------- APA 7

def initials(given):
    if not given:
        return ""
    given = clean(given)
    if ARABIC_RE.search(given):
        return given
    out = []
    # Splitting on dots too turns 'W.F.' into two initials.
    for part in re.split(r"[\s.]+", given):
        if not part:
            continue
        if re.fullmatch(r"[A-Z]{2,3}", part):          # Crossref often stores 'AJ'
            out += [c + "." for c in part]
        elif "-" in part:
            out.append("-".join(p[0].upper() + "." for p in part.split("-") if p))
        else:
            out.append(part[0].upper() + ".")
    return " ".join(out)


def apa_author(a):
    if a.get("org") or not a.get("given"):
        return a["family"]
    sep = "، " if ARABIC_RE.search(a["family"]) else ", "
    return f"{a['family']}{sep}{initials(a['given'])}"


def apa_authors(authors, arabic=False):
    names = [apa_author(a) for a in authors]
    if not names:
        return ""
    if len(names) == 1:
        return names[0]
    sep, last = ("، ", "، و") if arabic else (", ", ", & ")
    if len(names) <= 20:
        return sep.join(names[:-1]) + last + names[-1]
    return sep.join(names[:19]) + sep + "... " + names[-1]


def apa(rec):
    """APA 7 reference in markdown (*italics*). Titles are kept as published:
    APA wants sentence case, which needs a human eye for proper nouns.
    Arabic-script records use Arabic punctuation, as Arab universities expect."""
    authors = rec.get("authors") or []
    arabic = bool(ARABIC_RE.search(rec.get("title", "") + (authors[0]["family"] if authors else "")))
    c = "، " if arabic else ", "
    who = apa_authors(authors, arabic) or rec.get("title", "")
    if not who.endswith("."):
        who += "."                      # group authors: 'American Psychological Association.'
    year = rec.get("year") or ("د.ت" if arabic else "n.d.")
    title = rec.get("title", "").rstrip(".")
    parts = [f"{who} ({year})."]
    if rec.get("type") == "book":
        parts.append(f"*{title}*.")
        if rec.get("publisher"):
            parts.append(f"{rec['publisher']}.")
    elif rec.get("container"):
        parts.append(f"{title}.")
        src = f"*{rec['container']}*"
        real = lambda v: v if str(v or "").strip("0-") else ""          # in-press records carry 0 / 0-0
        if real(rec.get("volume")):
            src += f"{c}*{rec['volume']}*"
            if real(rec.get("issue")):
                src += f"({rec['issue']})"
        pages, art = real(rec.get("pages", "")), real(rec.get("article_number", ""))
        # Elsevier e-locators arrive as both page and article-number: APA wants 'Article N'.
        if art and (not pages or pages == art or "-" not in pages):
            src += f"{c}{'مقالة' if arabic else 'Article'} {art}"
        elif pages:
            src += f"{c}{pages}"
        parts.append(src + ".")
    else:
        parts.append(f"*{title}*.")
        if rec.get("publisher"):
            parts.append(f"{rec['publisher']}.")
    if rec.get("doi"):
        parts.append(f"https://doi.org/{rec['doi']}")
    elif rec.get("url"):
        parts.append(rec["url"])
    if (rec.get("retraction") or {}).get("status") == "retracted":
        parts.append("(مسحوب)" if arabic else "(Retracted)")
    return clean(" ".join(parts))


# ---------------------------------------------------------------- Zotero export (RIS, BibTeX, CSL-JSON)

RIS_TYPES = {"journal-article": "JOUR", "book": "BOOK", "book-chapter": "CHAP", "proceedings-article": "CONF",
             "dissertation": "THES", "report": "RPRT", "posted-content": "UNPB", "dataset": "DATA",
             "monograph": "BOOK", "edited-book": "EDBOOK", "reference-book": "BOOK"}
BIBTEX_TYPES = {"journal-article": "article", "book": "book", "book-chapter": "incollection",
                "proceedings-article": "inproceedings", "dissertation": "phdthesis", "report": "techreport",
                "monograph": "book", "edited-book": "book", "reference-book": "book"}
CSL_TYPES = {"journal-article": "article-journal", "book": "book", "book-chapter": "chapter",
             "proceedings-article": "paper-conference", "dissertation": "thesis", "report": "report",
             "posted-content": "article", "dataset": "dataset", "monograph": "book", "edited-book": "book",
             "reference-book": "book"}


def _pages(e):
    """(first, last) from '1237-1245'; an e-locator or a single page gives (x, '')."""
    pages = str(e.get("pages") or "").strip()
    if not pages.strip("0-"):
        return str(e.get("article_number") or "").strip(), ""
    bits = [p.strip() for p in pages.split("-", 1)]
    return bits[0], (bits[1] if len(bits) == 2 and bits[1] != bits[0] else "")


def _ident(e):
    """ISSN for articles, ISBN for books (Zotero reads both from SN / issn / isbn)."""
    if e.get("type") == "book":
        return str(e.get("isbn") or "")
    return ", ".join(hyphen_issn(x) for x in (e.get("issn") or []) if hyphen_issn(x))


def ris_entry(e):
    """One RIS record. Zotero maps T2 to the publication title and JO to the journal
    abbreviation, so both carry the container; AU is 'Family, Given' per author."""
    ty = RIS_TYPES.get(e.get("type") or "", "GEN")
    lines = [("TY", ty)]
    for a in e.get("authors") or []:
        lines.append(("AU", a["family"] if a.get("org") or not a.get("given") else f"{a['family']}, {a['given']}"))
    lines.append(("TI", clean(e.get("title"))))
    if e.get("container"):
        lines += [("T2", e["container"]), ("JO", e["container"])]
    if str(e.get("volume") or "").strip("0-"):
        lines.append(("VL", e["volume"]))
    if str(e.get("issue") or "").strip("0-"):
        lines.append(("IS", e["issue"]))
    sp, ep = _pages(e)
    if sp and ty != "BOOK":                       # SP on a BOOK record means 'number of pages' in Zotero
        lines.append(("SP", sp))
        if ep:
            lines.append(("EP", ep))
    if e.get("year"):
        lines.append(("PY", str(e["year"])))
    if e.get("doi"):
        lines += [("DO", e["doi"]), ("UR", "https://doi.org/" + e["doi"])]
    elif e.get("openlibrary"):
        lines.append(("UR", e["openlibrary"]))
    if _ident(e):
        lines.append(("SN", _ident(e)))
    if e.get("publisher"):
        lines.append(("PB", e["publisher"]))
    if e.get("abstract"):
        lines.append(("AB", clean(e["abstract"])))
    note = f"ledger key {e.get('key', '')}" + (f"; supports: {clean(e['note'])}" if e.get("note") else "")
    lines.append(("N1", note))
    return "\n".join(f"{tag}  - {clean(val)}" for tag, val in lines) + "\nER  - \n"


def bibtex_escape(s):
    s = clean(s)
    return re.sub(r"([&%$#_{}])", r"\\\1", s)


def bibtex_entry(e):
    """One BibTeX entry; the key is the ledger key, text stays UTF-8 (biblatex/biber
    and Zotero read it; classic bibtex8 needs an ASCII key)."""
    kind = BIBTEX_TYPES.get(e.get("type") or "", "misc")
    fields = []
    auths = [("{" + bibtex_escape(a["family"]) + "}") if a.get("org") or not a.get("given")
             else f"{bibtex_escape(a['family'])}, {bibtex_escape(a['given'])}" for a in e.get("authors") or []]
    if auths:
        fields.append(("author", " and ".join(auths)))
    fields.append(("title", bibtex_escape(e.get("title"))))
    if e.get("container"):
        fields.append(("booktitle" if kind in ("incollection", "inproceedings") else "journal", bibtex_escape(e["container"])))
    if str(e.get("volume") or "").strip("0-"):
        fields.append(("volume", bibtex_escape(e["volume"])))
    if str(e.get("issue") or "").strip("0-"):
        fields.append(("number", bibtex_escape(e["issue"])))
    sp, ep = _pages(e)
    if sp:
        fields.append(("pages", sp + ("--" + ep if ep else "")))
    if e.get("year"):
        fields.append(("year", str(e["year"])))
    if e.get("doi"):
        fields += [("doi", e["doi"]), ("url", "https://doi.org/" + e["doi"])]
    ident = _ident(e)
    if ident:
        fields.append(("isbn" if e.get("type") == "book" else "issn", ident))
    if e.get("publisher"):
        fields.append(("publisher", bibtex_escape(e["publisher"])))
    if e.get("abstract"):
        fields.append(("abstract", bibtex_escape(e["abstract"])))
    if e.get("note"):
        fields.append(("note", bibtex_escape(e["note"])))
    body = ",\n".join(f"  {k} = {{{v}}}" for k, v in fields)
    return f"@{kind}{{{e['key']},\n{body}\n}}\n"


def csl_entry(e):
    """CSL-JSON item (what Zotero, pandoc and citeproc read)."""
    authors = [{"literal": a["family"]} if a.get("org") else {"family": a["family"], "given": a.get("given", "")}
               for a in e.get("authors") or []]
    item = {"id": e["key"], "type": CSL_TYPES.get(e.get("type") or "", "article"), "title": clean(e.get("title"))}
    if authors:
        item["author"] = authors
    if e.get("container"):
        item["container-title"] = e["container"]
    if str(e.get("volume") or "").strip("0-"):
        item["volume"] = str(e["volume"])
    if str(e.get("issue") or "").strip("0-"):
        item["issue"] = str(e["issue"])
    sp, ep = _pages(e)
    if sp:
        item["page"] = sp + ("-" + ep if ep else "")
    if e.get("year"):
        y = str(e["year"]).strip()
        item["issued"] = {"date-parts": [[int(y)]]} if y.isdigit() else {"raw": y}
    if e.get("doi"):
        item["DOI"] = e["doi"]
        item["URL"] = "https://doi.org/" + e["doi"]
    ident = _ident(e)
    if ident:
        item["ISBN" if e.get("type") == "book" else "ISSN"] = ident
    if e.get("publisher"):
        item["publisher"] = e["publisher"]
    if e.get("abstract"):
        item["abstract"] = clean(e["abstract"])
    if e.get("note"):
        item["note"] = clean(e["note"])
    return item


def export_refs(entries, fmt):
    if fmt == "ris":
        return "\n".join(ris_entry(e) for e in entries)
    if fmt == "bibtex":
        return "\n".join(bibtex_entry(e) for e in entries)
    if fmt == "csl":
        return json.dumps([csl_entry(e) for e in entries], ensure_ascii=False, indent=1) + "\n"
    raise ValueError(fmt)


def csl_formatted(doi, style, locale=None):
    accept = f"text/x-bibliography; style={style}" + (f"; locale={locale}" if locale else "")
    r = http_get("https://doi.org/" + urllib.parse.quote(doi, safe="/()"), headers={"Accept": accept})
    # text/x-bibliography carries no charset; requests would guess ISO-8859-1.
    return clean(r.content.decode("utf-8")) if r is not None else ""


# ---------------------------------------------------------------- ledger

def load_ledger(path):
    if not os.path.exists(path):
        return []
    with open(path, encoding="utf-8") as f:
        return [json.loads(line) for line in f if line.strip()]


def save_ledger(path, entries):
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        for e in entries:
            f.write(json.dumps(e, ensure_ascii=False) + "\n")


def make_key(rec, existing):
    fam = rec["authors"][0]["family"] if rec.get("authors") else rec.get("title", "anon")
    base = (name_key(fam) or "anon") + str(rec.get("year") or "nd")
    key, n = base, 0
    while key in existing:
        n += 1
        key = base + "abcdefghijklmnopqrstuvwxyz"[n - 1]
    return key


def attach_abstract(rec):
    """Real abstract only. Missing stays missing: never paraphrase from memory."""
    doi = rec.get("doi")
    if not doi:
        return rec
    try:
        w = openalex_by_doi(doi)
    except HTTPError:
        w = None
    if w:
        rec["abstract"] = rebuild_abstract(w.get("abstract_inverted_index"))
        rec["abstract_source"] = "openalex" if rec["abstract"] else None
        rec["cited_by"] = {"openalex": w.get("cited_by_count"), "as_of": TODAY}
        if w.get("is_retracted") and rec["retraction"]["status"] == "ok":
            rec["retraction"]["status"] = "retracted"
            rec["retraction"]["notices"].append({"type": "retraction", "source": "openalex"})
    if not rec.get("abstract") and rec.get("crossref_abstract"):
        rec["abstract"] = re.sub(r"^(abstract|summary)\s*", "", rec["crossref_abstract"], flags=re.I)
        rec["abstract_source"] = "crossref"
    if not rec.get("abstract"):
        p = s2_by_doi(doi)
        if p and p.get("abstract"):
            rec["abstract"] = clean(p["abstract"])
            rec["abstract_source"] = "semanticscholar"
    if not rec.get("abstract"):
        ab = europepmc_abstract(doi)
        if ab:
            rec["abstract"], rec["abstract_source"] = ab, "europepmc"
    rec.pop("crossref_abstract", None)
    rec.setdefault("abstract", "")
    rec.setdefault("abstract_source", None)
    return rec


def europepmc_abstract(doi):
    """Europe PMC covers biomedical and life-science journals (free, no key)."""
    try:
        r = http_get("https://www.ebi.ac.uk/europepmc/webservices/rest/search",
                     params={"query": f'DOI:"{doi}"', "resultType": "core", "format": "json"}, retries=2)
    except HTTPError:
        return ""
    hits = (as_json(r).get("resultList") or {}).get("result", []) if r is not None else []
    return clean(hits[0].get("abstractText")) if hits else ""


# ---------------------------------------------------------------- manuscript parsing

YEAR = r"(?:19|20)\d{2}[a-z]?|n\.d\.|in press|د\.ت"
# Name particles in any case: 'Abu Orabi', 'Al Ajlouni', 'Bani Hani', 'van der Berg'.
SURNAME_EN = r"(?:(?i:de|van|von|der|den|da|del|al|el|bin|ibn|abu|bani|ben|bou|ould)[\s\-'])*[A-Z][\w'À-ɏ\-]+"
SURNAME_AR = r"[؀-ۿ]+(?:\s[؀-ۿ]+)?"


def strip_references(text):
    """Split body from the reference list (References / Bibliography / المراجع)."""
    m = re.search(r"^#{1,6}\s*(references|bibliography|works cited|reference list|المراجع|قائمة المراجع|المصادر والمراجع)\s*$",
                  text, re.I | re.M)
    if not m:
        return text, ""
    nxt = re.search(r"^#{1,6}\s+\S", text[m.end():], re.M)
    end = m.end() + nxt.start() if nxt else len(text)
    return text[: m.start()] + text[end:], text[m.end(): end]


YEAR_END = r"(?![\d/\-])"           # the 2026 in '2026-10-21' is a date, not a citation year
CITE_PREFIX = (r"^(?:see also|see|e\.g\.,?|cf\.|i\.e\.,?|as cited in|cited in|انظر|مثل|كما ورد في|كما في|راجع|مثلاً)"
               r"\s+")


def citation_author(chunk):
    """The author part of a citation ('Kirschner & Karpinski', 'Rosen et al.',
    'الأنصاري وآخرون'), or '' when the text before the year is not a citation."""
    c = re.sub(CITE_PREFIX, "", clean(chunk).strip(), flags=re.I).strip(" .")
    if not c:
        return ""
    if re.search(r"[A-Za-z]", c):
        m = re.search(SURNAME_EN, c)
        if not m or m.group(0).lower() in STOP_WORDS:
            return ""
        return c[m.start():].strip()            # 'كما ورد في Abdallah et al.' -> 'Abdallah et al.'
    words = re.findall(r"[؀-ۿ]+", c)
    while words and words[0] in AR_STOP:
        words.pop(0)                            # 'دراسة الخطيب' -> 'الخطيب'
    if re.search(r"\d", c) or len(words) > 5 or not words:
        return ""                               # '(37 من 352 استبانة، 2017)' is not a citation
    return " ".join(words)


def extract_citations(body):
    """Return [(author_chunk, year, snippet)] for APA author-date citations, EN and AR.

    author_chunk keeps the whole author part so ledger_match can tell
    'Junco (2012)' from 'Junco & Cotten (2012)'."""
    body = body.translate(ARABIC_DIGITS)
    found = []
    year_list = r"((?:" + YEAR + r")" + YEAR_END + r"(?:[,،]\s*(?:" + YEAR + r")" + YEAR_END + r")*)"
    # Parenthetical: (Kirschner & Karpinski, 2010; Junco, 2012a, p. 4) / (العتيبي، 2020)
    for m in re.finditer(r"\(([^()]*?(?:" + YEAR + r")[^()]*)\)", body):
        for part in re.split(r";|؛", m.group(1)):
            pm = re.search(r"^(.*?)[,،]\s*" + year_list, part.strip())
            if not pm:
                continue
            # The author is the last comma chunk before the year:
            # '(Monte Carlo method, Preacher & Selig, 2012)' -> 'Preacher & Selig'.
            chunks = [c.strip() for c in re.split(r"[,،]", pm.group(1)) if c.strip()]
            author = citation_author(chunks[-1] if chunks else "")
            if author:
                for y in re.split(r"[,،]\s*", pm.group(2)):
                    found.append((author, y.strip(), m.group(0)[:80]))
    # Narrative: Kirschner and Karpinski (2010) / Rosen et al. (2013) /
    # Liu, Kirschner & Karpinski (2017) / العتيبي (2020) / الأنصاري وآخرون (2026)
    narr_en = (r"(" + SURNAME_EN + r"(?:,\s*" + SURNAME_EN + r")*(?:,?\s+(?:and|&)\s+" + SURNAME_EN + r")?"
               r"(?:\s+et al\.?)?)\s+\((" + YEAR + r")" + YEAR_END + r"(?:[,،][^)]*)?\)")
    narr_ar = (r"(" + SURNAME_AR + r"(?:\s+(?:وآخرون|وزملاؤه|و[؀-ۿ]+))?)\s+\((" + YEAR + r")" + YEAR_END
               + r"(?:[,،][^)]*)?\)")
    for pat in (narr_en, narr_ar):
        for m in re.finditer(pat, body):
            author = citation_author(m.group(1))
            if author:
                found.append((author, m.group(2), m.group(0)[:80]))
    # De-duplicate identical (first author, year) found by both passes at one spot.
    seen, out = set(), []
    for au, yr, snip in found:
        k = (name_key(first_surname(au)), yr.rstrip("abcdefghijklmnopqrstuvwxyz"), snip)
        if k not in seen:
            seen.add(k)
            out.append((au, yr, snip))
    return out


STOP_WORDS = {"table", "figure", "fig", "chapter", "section", "appendix", "equation", "step", "phase",
              "model", "study", "hypothesis", "the", "in", "and", "as", "since", "from", "by", "of", "see",
              "january", "february", "march", "april", "may", "june", "july", "august", "september",
              "october", "november", "december", "spring", "summer", "fall", "autumn", "winter"}
AR_STOP = {"عام", "سنة", "في", "منذ", "جدول", "شكل", "الفصل", "الشكل", "الجدول", "دراسة", "حتى", "عن", "من"}


def first_surname(names):
    names = names.strip()
    if ARABIC_RE.search(names[:30]):
        words = re.findall(r"[؀-ۿ]+", names)
        words = [w for w in words if w not in AR_STOP and w not in ("وآخرون", "وزملاؤه")]
        return words[0] if words else ""
    m = re.search(SURNAME_EN, names)
    if not m:
        return ""
    w = m.group(0)
    return "" if w.lower() in STOP_WORDS else w


def parse_authors_chunk(chunk):
    """'Liu, Kirschner & Karpinski' -> ('Liu', ['Kirschner', 'Karpinski'], 'multi');
    'Rosen et al.' -> ('Rosen', [], 'etal'); 'بدر وحميد' -> ('بدر', ['حميد'], 'multi')."""
    c = clean(chunk)
    etal = bool(re.search(r"\bet al\b|وآخرون|وزملاؤه", c))
    c = re.sub(r"\bet al\.?|وآخرون|وزملاؤه", " ", c)
    if re.search(r"[A-Za-z]", c):
        parts = re.split(r"\s*,\s*|\s*&\s*|\s+and\s+", c)
    else:
        parts = re.split(r"[,،]|\s+و(?=[؀-ۿ])", c)
    names = [first_surname(p.strip(" .")) or p.strip(" .") for p in parts if p.strip(" .")]
    if not names:
        return "", [], "single"
    return names[0], names[1:], "etal" if etal else ("multi" if len(names) > 1 else "single")


def surname_eq(a, b):
    a, b = name_key(a), name_key(b)
    return bool(a) and bool(b) and (a == b or a.endswith(b) or (b.endswith(a) and len(a) > 3))


def ledger_match(sur, year, ledger):
    """Ledger entries cited as `sur (year)`; sur may be a bare surname or the
    whole author chunk from extract_citations."""
    y = re.sub(r"[a-z]$", "", year)
    first, others, form = parse_authors_chunk(sur)
    if not first:
        return []
    hits = []
    for e in ledger:
        fams = [a["family"] for a in e.get("authors") or []]
        if not fams:
            continue
        ey = str(e.get("year") or "n.d.")
        if surname_eq(fams[0], first) and (ey == y or y in ("n.d.", "in press", "د.ت")):
            hits.append(e)
    if len(hits) <= 1:
        return hits
    # Several papers by the same first author and year: APA 7 tells them apart
    # by the co-authors ('Junco (2012)' vs 'Junco & Cotten (2012)').
    def fits(e):
        fams = [a["family"] for a in e["authors"]]
        if form == "etal":
            return len(fams) >= 3
        if form == "multi":
            return len(fams) == len(others) + 1 and all(surname_eq(f, o) for f, o in zip(fams[1:], others))
        return len(fams) == 1
    strict = [e for e in hits if fits(e)]
    return strict or hits


# ---------------------------------------------------------------- numbers check

# Statistic labels, as written in papers and in analysis output, mapped to one name.
LABEL_ALIASES = {
    "r": "r", "pearson": "r", "rho": "rho", "ρ": "rho", "spearman": "rho", "rs": "rho",
    "df1": "df", "df2": "df", "df_between": "df", "df_within": "df", "median": "median", "mdn": "median",
    "vif": "vif", "skew": "skew", "kurtosis": "kurtosis", "min": "min", "max": "max",
    "w": "w", "shapiro": "w", "u": "u", "mann_whitney": "u", "h": "h", "kruskal": "h",
    "sr2": "sr2", "sr²": "sr2", "cooks_d": "cooks_d", "cook": "cooks_d", "diff": "diff", "difference": "diff",
    "mean_diff": "diff", "levene": "levene", "bp": "bp", "breusch_pagan": "bp",
    # power, SEM and PLS labels written by effects.py and analyze_sem.py
    "f2": "f2", "f²": "f2", "cohens_f": "cohen_f", "g": "g", "hedges_g": "g", "power": "power",
    "cr": "cr", "ave": "ave", "htmt": "htmt", "cfi": "cfi", "tli": "tli", "rmsea": "rmsea", "srmr": "srmr",
    "loading": "loading", "pm": "pm", "indirect": "indirect", "path_a": "a", "path_b": "b", "c_prime": "c_prime",
    "a": "a", "c": "c", "total": "c", "ab": "indirect", "indirect_b": "indirect", "mediated": "pm",
    "proportion_mediated": "pm",
    "kappa": "kappa", "κ": "kappa",
    # Arabic statistic names as written in Arabic theses
    "ن": "n", "بيتا": "beta", "ألفا": "alpha", "الفا": "alpha", "ت": "t", "ف": "f", "ر": "r", "كا²": "chi2", "كا2": "chi2",
    "β": "beta", "beta": "beta", "std_beta": "beta", "stdbeta": "beta", "standardized": "beta",
    "b": "b", "coef": "b", "coefficient": "b", "estimate": "b", "unstandardized": "b",
    "se": "se", "stderr": "se", "std_err": "se", "t": "t", "z": "z", "f": "f",
    "p": "p", "sig": "p", "pvalue": "p", "p_value": "p",
    "r2": "r2", "r²": "r2", "rsquared": "r2", "r_squared": "r2", "adj_r2": "adj_r2", "adjr2": "adj_r2",
    "m": "m", "mean": "m", "avg": "m", "average": "m", "المتوسط": "m", "متوسط": "m",
    "sd": "sd", "std": "sd", "stdev": "sd", "الانحراف": "sd", "المعياري": "sd",
    "n": "n", "count": "n", "df": "df", "α": "alpha", "alpha": "alpha", "cronbach": "alpha",
    "d": "d", "cohen": "d", "cohens_d": "d", "ci": "ci", "lower": "ci", "upper": "ci", "ll": "ci", "ul": "ci",
    "ci_low": "ci", "ci_high": "ci", "or": "or", "χ2": "chi2", "χ²": "chi2", "chi2": "chi2", "chisq": "chi2",
    "η2": "eta2", "η²": "eta2", "eta2": "eta2", "eta_sq": "eta2", "ηp2": "eta2", "pct": "pct", "percent": "pct",
}
SIGN_FREE = {"t", "z", "f", "chi2", "p", "n", "df", "sd", "se", "r2", "adj_r2", "eta2", "alpha", "pct"}
LABEL_WORD = re.compile(r"([A-Za-z_αβρχηω²؀-ۿ][\w.²]*)\s*(?:\([^()]{0,12}\))?\s*([=:<>≤≥])?\s*$")
NUM_TOKEN = re.compile(r"(?<![\w.])([<>≤≥=]\s*)?(-?(?:\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+\.\d+|\.\d+|\d+))(\s*%)?(?![\w.]*\d)")
NOT_STATS = re.compile(r"\b(?:section|sec\.|chapter|ch\.|subsection|part)\s*\d+(?:\.\d+)*"
                       r"|(?:القسم|الفصل|الفقرة|البند|الجزء|المبحث|المطلب)\s*(?:رقم\s*)?\d+(?:\.\d+)*"
                       r"|\b(?:table|tables|figure|fig\.|model|step|hypothesis|hypotheses|study|studies|"
                       r"appendix|item|items|wave|phase|stage|question)\s*\d+[a-z]?(?![.\d])"
                       r"|\b(?:H|RQ)\d+[a-z]?\b"
                       r"|(?:جدول|الجدول|شكل|الشكل|الفرضية|الفرضيات|السؤال|الفصل|الملحق|المحور|البعد)\s*(?:رقم\s*)?\d+(?![.\d])",
                       re.I)
SAMPLE_NOUN = re.compile(r"^\s+(?:[A-Za-z\-]+\s+){0,2}(students|participants|respondents|adolescents|adults|employees|"
                         r"teachers|nurses|patients|children|undergraduates|workers|users|subjects|managers|"
                         r"طالب|طالبة|طالبًا|طالباً|طلاب|مشارك|مشاركًا|مشاركاً|موظف|موظفًا|موظفاً|معلم|معلمًا|معلماً|ممرض|مريض|فرد|فردًا|فرداً|مستجيب)",
                         re.I)


def canon_label(word):
    if not word:
        return None
    w = word.strip().strip("_.").replace("'", "")
    segs = re.split(r"[._~]", w)
    last, last2 = segs[-1], "_".join(segs[-2:])
    for cand in (w, last2, last, w.lower(), last2.lower(), last.lower()):
        if cand in LABEL_ALIASES:
            return LABEL_ALIASES[cand]
    if any(s.lower() == "ci" for s in segs):
        return "ci"
    return None


DECLARATION_HEAD = re.compile(r"^#{1,6}\s*(declaration of generative ai|generative ai use|use of (generative )?ai|"
                              r"ai[- ]use (statement|disclosure)|الإفصاح عن استخدام الذكاء الاصطناعي)", re.I)


def table_labels(lines):
    """Pipe-table cells inherit the statistic named in their column header (M, SD, r, B...)."""
    labels, header = {}, None
    for ln, line in enumerate(lines, 1):
        if not line.strip().startswith("|"):
            header = None
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
            continue
        if header is None:
            header = [canon_label(re.sub(r"[*_`]", "", c).split("(")[0].strip()) for c in cells]
            continue
        labels[ln] = header
    return labels


def numbers_in_text(text, include_ints=False):
    text = re.sub(r"<!--.*?-->", lambda m: "\n" * m.group(0).count("\n"), text, flags=re.S)
    # Volumes, issues and pages in the reference list are never results.
    body, refsec = strip_references(text)
    if refsec:
        text = body + "\n" * refsec.count("\n")
    text = text.translate(ARABIC_DIGITS).translate(DASHES)
    # Drop the AI-use declaration (model versions are not results) up to the next heading.
    kept, skipping = [], False
    for line in text.splitlines():
        if re.match(r"^#{1,6}\s", line):
            skipping = bool(DECLARATION_HEAD.match(line))
        kept.append("" if skipping else line)
    text = "\n".join(kept)
    text = re.sub(r"(?<![\w*])[*_]{1,2}(?=[^\s*_])|(?<=[^\s*_])[*_]{1,2}(?![\w*])", "", text)   # *r*(197) -> r(197)
    cell_labels = table_labels(text.splitlines())
    text = re.sub(r"https?://\S+|10\.\d{4,9}/\S+", " ", text)          # DOIs, URLs
    text = re.sub(r"\((?:[^()]*?(?:19|20)\d{2}[a-z]?)\)", " ", text)   # citations
    text = re.sub(r"\[[^\]]*[A-Za-z؀-ۿ][^\]]*\]", " ", text)   # [placeholders with words]; CI brackets stay
    text = NOT_STATS.sub(" ", text)
    text = re.sub(r"أقل من|اقل من|دون", "<", text)
    text = re.sub(r"أكبر من|اكبر من|أعلى من|يزيد عن", ">", text)
    # Conventional confidence levels are not results.
    text = re.sub(r"\b(90|95|99)\s*%\s*(?=CI\b|confidence|فترة|مجال)", " ", text, flags=re.I)
    text = re.sub(r"(فترة|مجال) (ال)?ثقة\s*(90|95|99)\s*%", " ", text)
    out = []
    for ln, line in enumerate(text.splitlines(), 1):
        if re.match(r"\s*#", line):
            continue
        col_of = None
        if ln in cell_labels:
            bounds = [i for i, ch in enumerate(line) if ch == "|"]
            col_of = lambda pos: next((k - 1 for k, b in enumerate(bounds) if b > pos), None)
        for m in NUM_TOKEN.finditer(line):
            op, val, pct = m.group(1), m.group(2), m.group(3)
            before = line[max(0, m.start() - 30): m.start()]
            lw = LABEL_WORD.search(before)
            word = lw.group(1) if lw else None
            # A lone letter is a label only with an operator ('a = .48'), so
            # 'a 5-point scale' is not read as path a. The operator may sit in
            # the number token (op) or just before it (lw.group(2)).
            if word and len(word) == 1 and word.isascii() and not (op or lw.group(2)):
                word = None
            label = "pct" if pct else canon_label(word)
            if re.search(r"cook'?s?\s*D?\s*[=:]?\s*$", before, re.I):
                label = "cooks_d"
            val = val.replace(",", "")                                 # 15,199 -> 15199
            if col_of and label is None:
                c = col_of(m.start())
                hdr = cell_labels[ln]
                label = hdr[c] if c is not None and 0 <= c < len(hdr) else None
            if re.search(r"(?:\b(?:t|F|r|rs|z)|χ2|χ²)\s*\(\s*$|(?:\bF|χ2|χ²)\s*\(\s*\d+(?:\.\d+)?\s*,\s*$", before):
                label = "df"            # t(196), F(2, 196): degrees of freedom
            elif label is None and re.search(r"\bCI\b|\[[^\]]*$|فترة الثقة", line[max(0, m.start() - 15): m.start()]):
                label = "ci"
            if "." not in val and not op and not pct:
                sample = SAMPLE_NOUN.match(line[m.end():])
                if sample and int(val) >= 10:
                    label = "n"
                elif include_ints and int(val) >= 10 and not re.fullmatch(r"(19|20)\d{2}", val):
                    pass                 # literature reviews: every count quoted from a source is a claim
                elif label is None or re.fullmatch(r"(19|20)\d{2}", val) or int(val) < 2:
                    continue
            out.append({"line": ln, "op": (op or "").strip(), "value": val, "pct": bool(pct), "label": label,
                        "context": line.strip()[:120]})
    return out


def evidence_numbers(paths=(), texts=()):
    """All numbers in the evidence, each with the label word just before it."""
    nums = []
    chunks = []
    for p in paths:
        with open(p, encoding="utf-8", errors="replace") as f:
            chunks.append(f.read())
    chunks += list(texts)
    for t in chunks:
        t = t.translate(ARABIC_DIGITS)
        for m in re.finditer(r"(?<![A-Za-z_\d.])-?(?:\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+\.\d+(?:[eE][-+]?\d+)?|\.\d+|\d+(?:[eE][-+]?\d+)?)", t):
            try:
                v = float(m.group(0).replace(",", ""))
            except ValueError:
                continue
            lw = LABEL_WORD.search(t[max(0, m.start() - 40): m.start()])
            w = lw.group(1) if lw else None
            if w and len(w) == 1 and w.isascii() and not lw.group(2):
                w = None
            label = canon_label(w)
            if label is None and t[m.end(): m.end() + 2].strip().startswith("%"):
                label = "pct"
            nums.append((label, v))
    return nums


def number_supported(item, evidence):
    """Return 'strict', 'weak', 'sign' (only |value| matches) or None.

    strict: matched a value carrying the same label in the evidence.
    weak:   the evidence has no value with that label, so any value was allowed;
            label the analysis output (r = 0.312) to make the check strict.
    """
    v = float(item["value"])
    dec = len(item["value"].split(".")[1]) if "." in item["value"] else 0
    tol = 0.5 * 10 ** (-dec) + 1e-12
    lab = item["label"]
    labelled = [e for l, e in evidence if l == lab] if lab else []
    if lab == "pct":
        # '13% of the variance' is r2 = 0.13 in the evidence.
        labelled = [e for l, e in evidence if l in ("pct", None, "r2", "adj_r2", "eta2", "sr2", "pm")]
    pool, kind = (labelled, "strict") if labelled else ([e for _, e in evidence], "weak")
    if item["op"] in ("<", "≤"):
        return kind if any(0 <= e < v or (lab != "p" and abs(e) < v) for e in pool) else None
    if item["op"] in (">", "≥"):
        return kind if any(abs(e) > v for e in pool) else None
    cands = [(v, tol)] + ([(v / 100, tol / 100)] if item["pct"] else [])
    for c, t in cands:
        if any(abs(e - c) <= t for e in pool):
            return kind
    for c, t in cands:
        if any(abs(abs(e) - abs(c)) <= t for e in pool):
            return kind if (lab in SIGN_FREE or lab is None and kind == "weak") else "sign"
    return None


# ---------------------------------------------------------------- SJR

def load_sjr(path):
    if not path or not os.path.exists(path):
        sys.exit("SJR file not found. Download it in your own browser from "
                 "https://www.scimagojr.com/journalrank.php (button 'Download data'), "
                 "then pass --sjr <file> or set SJR_CSV. Scripts cannot fetch it: "
                 "the site blocks bots and its robots.txt disallows the export.")
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f, delimiter=";"))
    head = rows[0]
    idx = {}
    for i, h in enumerate(head):
        idx.setdefault(h.strip(), i)              # keeps the FIRST 'Publisher'

    def col(*names, regex=None):
        for n in names:
            if n in idx:
                return idx[n]
        if regex:
            for h, i in idx.items():
                if re.search(regex, h):
                    return i
        return None

    c = {
        "title": col("Title"), "issn": col("Issn"), "sjr": col("SJR"),
        "q": col("SJR Best Quartile", "SJR Quartile"), "h": col("H index"),
        "cats": col("Categories"), "areas": col("Areas"), "type": col("Type"),
        "pub": col("Publisher"), "country": col("Country"), "region": col("Region"),
        "oa": col("Open Access"), "coverage": col("Coverage"), "rank": col("Rank"),
        "sourceid": col("Sourceid"),
    }
    edition = None
    for h in head:
        m = re.search(r"Total Docs\. \((\d{4})\)", h)
        if m:
            edition = m.group(1)
    m = re.search(r"(19|20)\d{2}", os.path.basename(path))
    edition = edition or (m.group(0) if m else "unknown")
    journals = []
    for r in rows[1:]:
        if len(r) < len(head) - 1:
            continue
        g = lambda k: r[c[k]].strip() if c[k] is not None and c[k] < len(r) else ""
        cats = re.findall(r"(.+?) \((Q[1-4])\)(?:; |$)", g("cats"))
        journals.append({
            "title": g("title"),
            "issns": [compact_issn(x) for x in g("issn").split(",") if compact_issn(x)],
            "sjr": float(g("sjr").replace(",", ".")) if g("sjr") else None,
            "best_q": g("q") if g("q") in ("Q1", "Q2", "Q3", "Q4") else "not ranked",
            "h": g("h"),
            "categories": [(n.strip(), q) for n, q in cats] or [(x.strip(), "-") for x in g("cats").split(";") if x.strip()],
            "areas": g("areas"), "type": g("type"), "publisher": g("pub"), "country": g("country"),
            "region": g("region"), "oa": g("oa"), "coverage": g("coverage"), "sourceid": g("sourceid"),
        })
    return journals, edition


# ---------------------------------------------------------------- commands

def fetch_candidates(q, a):
    """One query against the chosen source(s) -> (rows, total)."""
    rows, total = [], None
    if a.source in ("openalex", "both"):
        rows, total = search_openalex(q, a)
    if a.source == "s2" or (a.source == "both"):
        try:
            more, t2 = search_s2(q, a)
        except HTTPError as e:
            print(f"[semantic scholar] skipped: {e}", file=sys.stderr)
            more, t2 = [], None
        have = {r["doi"].lower() for r in rows if r["doi"]}
        rows += [r for r in more if not r["doi"] or r["doi"].lower() not in have]
        total = total or t2
    if a.source == "crossref":
        rows, total = search_crossref(q, a)
    return rows, total


def cmd_search(a):
    if a.show_expansions:
        print("Arabic synonym map used by --expand (a term with no variants is kept as is):")
        print(table([[t, " | ".join(v) or "(kept)"] for t, v in ARABIC_SYNONYMS.items()], ["term", "variants"]))
        print(f"At most {MAX_EXPANSIONS} variants run per search, one substitution each; every variant is one more "
              "API call (10 OpenAlex credits when the source is OpenAlex).")
        return 0
    if a.issn and a.source not in ("openalex", "crossref"):
        print("--issn works with OpenAlex or Crossref only; using --source openalex", file=sys.stderr)
        a.source = "openalex"
    if a.out:
        a.abstracts = True            # saved candidates feed gapmap, effects and comap, which need abstracts
    rows, total = fetch_candidates(a.query, a)
    expansions = []
    if a.expand:
        variants, kept = expand_query(a.query)
        have = {r["doi"].lower() for r in rows if r["doi"]} | {journal_key(r["title"]) for r in rows}
        for term, variant, q2 in variants:
            more, _ = fetch_candidates(q2, a)
            new = [r for r in more if (r["doi"].lower() if r["doi"] else journal_key(r["title"])) not in have]
            for r in new:
                have.add(r["doi"].lower() if r["doi"] else journal_key(r["title"]))
            rows += new
            expansions.append((term, variant, q2, len(more), len(new)))
        if not variants and not kept:
            print("[expand] no term of the synonym map occurs in this query: only the original ran "
                  "(`search --show-expansions` lists the map)", file=sys.stderr)
        for term in kept:
            print(f"[expand] '{term}' kept as is (standard wording, no variant in the map)", file=sys.stderr)
    # The same paper often appears twice (online-first DOI and final DOI): keep one row.
    # Titles are compared loosely: 'organisational' / 'organizational' and small edits
    # between the two versions should still count as the same paper.
    uniq = []
    for r in rows:
        k = journal_key(r["title"]).replace("isation", "ization").replace("our", "or")
        twin = next((u for u in uniq if k and difflib.SequenceMatcher(None, k, u["_k"]).ratio() >= 0.93), None)
        if twin:
            twin.setdefault("also", []).append(r["doi"] or "(no DOI)")
            continue
        r["_k"] = k
        uniq.append(r)
    for r in uniq:
        r.pop("_k", None)
    rows = uniq
    print(f"query: {a.query or '(none)'} | source: {a.source} | total matches: {total} | shown: {len(rows)} | {TODAY}")
    if a.source == "crossref":
        print("Crossref: 'total matches' counts every work sharing a token with the query; relevance ranking is "
              "Crossref's, so judge the rows, not the count. No language, OA or core-journal filter exists there.")
    if expansions:
        print("expansions run: " + "; ".join(f"'{t}' -> '{v}' ({n_new} new of {n} hits)" for t, v, _, n, n_new in expansions))
    print(table([[str(i + 1), str(r["year"] or ""), str(r["cites"] if r["cites"] is not None else ""),
                  short(r["first_author"], 22) + (" et al." if r["n_authors"] > 2 else ""),
                  short(r["title"], 80),
                  short(r["venue"], 32) + (" *" if r.get("core_venue") is False and r["venue"] else ""),
                  r["doi"] or "(no DOI)"]
                 for i, r in enumerate(rows)],
                ["#", "year", "cites", "first author", "title", "venue", "doi"]))
    if any(r.get("core_venue") is False and r["venue"] for r in rows):
        print("* venue is outside the CWTS core list: vet it (`journal <ISSN>`) before relying on it, or rerun with --core.")
    for r in rows:
        if r.get("also"):
            print(f"duplicate: {r['doi']} also listed as {', '.join(r['also'])} (online-first vs final version: cite the final one)")
    if a.issn and "," in a.issn:
        counts = {}
        for r in rows:
            counts[r["venue"]] = counts.get(r["venue"], 0) + 1
        print("hits per journal (scope evidence): " + "; ".join(f"{v}: {n}" for v, n in
                                                             sorted(counts.items(), key=lambda x: -x[1])))
        print("journals with 0 hits above publish nothing matching this query in the period searched.")
    if a.out:
        with open(a.out, "w", encoding="utf-8") as f:
            json.dump({"query": a.query, "date": TODAY, "source": a.source, "results": rows}, f,
                      ensure_ascii=False, indent=1)
        print(f"saved {len(rows)} candidates to {a.out}")
    print("Candidates are NOT sources yet: read the abstract, then `add` the ones you will cite.")


def cmd_verify(a):
    bad = 0
    if a.citations_file:
        return verify_list(a)
    if a.citation:
        res, _ = match_citation(a.citation)
        print(table([[r["verdict"], str(r["title_similarity"]), "yes" if r["author_match"] else "no",
                      {True: "yes", False: "no", None: "?"}[r["year_match"]], str(r["year"]),
                      short(r["first_author"], 18), short(r["title"], 60), r["doi"], r["retraction"],
                      short("; ".join(r["discrepancies"]), 60)]
                     for r in res],
                    ["verdict", "title sim", "author", "year", "pub year", "first author", "title", "doi",
                     "retraction", "citation errors"]))
        verdict, best, detail = judge_citation(a.citation) if not any(
            r["verdict"] in ("match", "match-with-errors") for r in res) else (None, res[0], "")
        if verdict:
            print(f"{verdict}: {detail}. Do not cite it until a DOI or ISBN is confirmed.")
            bad = 1
        else:
            if best["verdict"] == "match-with-errors":
                print("REAL PAPER, the citation has errors: " + "; ".join(best["discrepancies"]))
            print("Corrected: " + apa(best["record"]))
            if best["retraction"] in ("retracted", "expression_of_concern"):
                print(f"WARNING: {best['retraction']}")
                bad = 1
        return bad
    rows = []
    for d in a.dois:
        doi = norm_doi(d)
        rec, verdict = resolve_doi(doi)
        if rec is None:
            rows.append([doi, verdict.upper(), "", "", "", ""])
            bad = 1
            continue
        st = rec["retraction"]["status"]
        if st in ("retracted", "expression_of_concern") or rec["retraction"]["is_update_notice"]:
            bad = 1
        flag = st + (" (this DOI IS a notice)" if rec["retraction"]["is_update_notice"] else "")
        if st in ("retracted", "expression_of_concern"):
            flag = notice_line(rec)
        fam = rec["authors"][0]["family"] if rec["authors"] else ""
        rows.append([doi, verdict, flag, str(rec["year"]), short(fam, 20), short(rec["title"], 70)])
    print(table(rows, ["doi", "exists", "integrity", "year", "first author", "title"]))
    return bad


def notice_line(rec):
    """'retracted 2010 by notice 10.1016/... (source: publisher, retraction-watch)'."""
    ns = [n for n in (rec.get("retraction") or {}).get("notices", []) if n.get("type") in INTEGRITY_FLAGS]
    if not ns:
        return (rec.get("retraction") or {}).get("status", "")
    first_n = ns[0]
    srcs = ", ".join(sorted({n.get("source") or "?" for n in ns}))
    return (f"{first_n['type'].replace('_', ' ')} {first_n.get('year') or ''} by notice "
            f"https://doi.org/{first_n.get('notice_doi')} (source: {srcs})").replace("  ", " ")


def split_reference_list(text):
    """One reference per line, or per blank-line block; strips '1.' / '[1]' numbering."""
    blocks = re.split(r"\n\s*\n", text.strip()) if re.search(r"\n\s*\n", text.strip()) else text.strip().splitlines()
    refs = []
    for b in blocks:
        b = re.sub(r"\s+", " ", b).strip()
        b = re.sub(r"^(\[\d+\]|\d+[.)])\s*", "", b)
        if len(b) > 25:
            refs.append(b)
    return refs


def verify_list(a):
    with open(a.citations_file, encoding="utf-8") as f:
        refs = split_reference_list(f.read())
    rows, found, bad, details = [], [], 0, []
    for i, ref in enumerate(refs, 1):
        dois = re.findall(r"10\.\d{4,9}/[^\s<>\"]+", ref)
        if dois:
            rec, v = resolve_doi(norm_doi(dois[0]))
            if rec is None:
                verdict, best, detail = "FABRICATED?", None, f"DOI {dois[0]} is not registered"
            else:
                diffs = discrepancies(parse_citation(ref), rec)
                sim = title_sim(rec["title"], parse_citation(ref).get("title", ref))
                if sim < 0.6:
                    verdict, detail = "WRONG DOI", f"DOI points to '{short(rec['title'], 60)}'"
                else:
                    verdict, detail = ("REAL" if not diffs else "REAL, ERRORS"), "; ".join(diffs)
                st = rec["retraction"]["status"]
                if st in ("retracted", "expression_of_concern"):
                    verdict, detail = st.upper(), st
                best = {"record": rec, "doi": rec["doi"]}
        else:
            verdict, best, detail = judge_citation(ref)
        if verdict.startswith("REAL"):
            found.append((i, best["record"]))
        else:
            bad = 1
        if verdict in ("RETRACTED", "CONCERN") and best:
            details.append(f"#{i}: " + notice_line(best["record"]))
        elif detail and not verdict == "REAL":
            details.append(f"#{i}: {verdict}: {detail}")
        rows.append([str(i), verdict, short(ref, 55), best["doi"] if best else "", short(detail, 70)])
    print(table(rows, ["#", "verdict", "as cited", "doi", "detail"]))
    details[:] = [d for d in details if d]
    if details:
        print("\nDetails:")
        for d in details:
            print("  " + d)
    print("\nCorrected references (from the registry, not from the list):\n")
    for i, rec in found:
        print(f"{i}. {apa(rec)}")
    found = [rec["doi"] for _, rec in found]
    print("\nREAL = exists and matches; REAL, ERRORS = exists but the citation has wrong details (use the corrected "
          "version); RETRACTED = remove or discuss the retraction; WRONG DOI / FABRICATED? / UNVERIFIED = do not "
          "cite until the user shows the source.")
    if a.add_to and found:
        cmd_add(argparse.Namespace(dois=found, isbn=[], ledger=a.add_to, note="from verified reference list",
                                   read="metadata", allow_retracted=False, edition_year=None, publisher=None))
    return bad


def add_manual(a):
    """A source with no DOI or ISBN (a scale manual, a report, a standard, a thesis):
    kept in the ledger but marked unverified, so the agent confirms it with the user."""
    if not (a.title and a.year):
        sys.exit("--manual needs --title and --year (plus --authors \"Family, Given; Family, Given\", --publisher, --url)")
    ledger = load_ledger(a.ledger)
    keys = {e["key"] for e in ledger}
    authors = []
    for part in (a.authors or "").split(";"):
        part = part.strip()
        if part:
            fam, _, giv = part.partition(",")
            authors.append({"family": clean(fam), "given": clean(giv)})
    rec = {"type": a.type, "doi": None, "isbn": None, "title": clean(a.title), "authors": authors, "year": a.year,
           "container": "", "volume": "", "issue": "", "pages": "", "article_number": "", "issn": [],
           "publisher": clean(a.publisher or ""), "url": (a.url or "").strip(), "abstract": "", "abstract_source": None,
           "retraction": {"status": "n/a", "notices": [], "is_update_notice": False},
           "verified": {"by": "manual", "at": TODAY}, "note": a.note or "", "read_level": a.read, "added": TODAY}
    rec["key"] = make_key(rec, keys)
    rec["apa"] = apa(rec)
    ledger.append(rec)
    save_ledger(a.ledger, ledger)
    print(f"added {rec['key']} (MANUAL, unverified: no DOI or ISBN; confirm every detail with the user): {rec['apa']}")
    print(f"1 added | ledger: {a.ledger} ({len(ledger)} sources)")


def cmd_add(a):
    if a.dois_file:
        with open(a.dois_file, encoding="utf-8") as f:
            a.dois += [d.strip() for d in f if d.strip() and not d.startswith("#")]
    if a.manual:
        return add_manual(a)
    ledger = load_ledger(a.ledger)
    keys = {e["key"] for e in ledger}
    have = {(e.get("doi") or "").lower() for e in ledger} | {(e.get("isbn") or "") for e in ledger}
    added = 0
    targets = [("doi", norm_doi(d)) for d in a.dois] + [("isbn", i) for i in a.isbn]
    for kind, ident in targets:
        if ident.lower() in have or re.sub(r"[^0-9Xx]", "", ident) in have:
            print(f"skip {ident}: already in ledger")
            continue
        if kind == "doi":
            rec, verdict = resolve_doi(ident)
            if rec is None:
                print(f"REFUSED {ident}: {verdict}. A citation that cannot be resolved does not enter the ledger.")
                continue
            if rec["retraction"]["is_update_notice"]:
                print(f"REFUSED {ident}: this DOI is a retraction/correction NOTICE, not the paper.")
                continue
            if rec["retraction"]["status"] in ("retracted", "expression_of_concern") and not a.allow_retracted:
                print(f"REFUSED {ident}: {rec['retraction']['status']}. Use --allow-retracted only when the text "
                      "discusses the retraction itself, and say so in the sentence.")
                continue
            rec = attach_abstract(rec)
        else:
            hits = book_lookup(isbn=ident)
            if not hits:
                print(f"REFUSED ISBN {ident}: not found in Open Library. Check the ISBN on the book itself.")
                continue
            rec = hits[0]
            rec["abstract"], rec["abstract_source"] = "", None
            if a.edition_year:
                rec["year"] = a.edition_year
            if a.publisher:
                rec["publisher"] = a.publisher
        rec["key"] = make_key(rec, keys)
        keys.add(rec["key"])
        rec["note"] = a.note or ""
        rec["read_level"] = a.read
        rec["added"] = TODAY
        rec["apa"] = apa(rec)
        ledger.append(rec)
        have.add((rec.get("doi") or rec.get("isbn") or "").lower())
        added += 1
        warn = "" if rec.get("abstract") else "  [no abstract found: read the paper before citing a specific finding]"
        print(f"added {rec['key']}: {rec['apa']}{warn}")
        for hint in record_hints(rec):
            print(f"  check: {hint}")
    if added:
        save_ledger(a.ledger, ledger)
    print(f"{added} added | ledger: {a.ledger} ({len(ledger)} sources)")


def looks_title_case(title):
    words = [w for w in re.findall(r"[A-Za-z][A-Za-z'\-]*", title)[1:] if len(w) > 3]
    return len(words) >= 3 and sum(w[0].isupper() for w in words) / len(words) > 0.6


def sentence_case(title):
    """APA sentence case suggestion: first word and first word after a colon keep
    their capital, acronyms (all caps, digits) stay, everything else is lowered,
    including the second part of hyphenated words. Proper nouns (Jordan, Facebook)
    still need a human eye, so the hint says so."""
    out, cap = [], True
    for w in re.split(r"(\s+)", title.strip().rstrip(".")):
        if not w.strip():
            out.append(w)
            continue
        parts = w.split("-")
        fixed = []
        for k, p in enumerate(parts):
            letters = re.sub(r"[^A-Za-z]", "", p)
            acronym = (len(letters) >= 2 and letters.isupper()) or re.search(r"\d", p)
            if acronym or (cap and k == 0):
                fixed.append(p)
            else:
                fixed.append(p[:1].lower() + p[1:] if p[1:2].islower() or len(p) == 1 else p.lower())
        out.append("-".join(fixed))
        cap = w.endswith((":", "?", "!"))
    return "".join(out)


def record_hints(rec):
    hints = []
    t = rec.get("title", "")
    if t and not ARABIC_RE.search(t) and looks_title_case(t):
        hints.append(f"Title Case: APA wants sentence case. Suggested, after you re-capitalise proper nouns "
                     f"(countries, platforms, theories named after people): fix {rec['key']} --title \"{sentence_case(t)}\"")
    for i, au in enumerate(rec.get("authors") or [], 1):
        if au.get("org") or ARABIC_RE.search(au["family"]) or i in (rec.get("confirmed_authors") or []):
            continue
        if (" " in au["family"] and not re.match(r"(?i)(al|el|de|van|von|da|abu|bin|ibn)[\s\-]", au["family"])) \
                or not au.get("given"):
            hints.append(f"author {i} '{au['family']}, {au.get('given', '')}' may be split wrongly: "
                         f"fix {rec['key']} --author \"{i}:Family,Given\"")
    ys = rec.get("years_seen") or []
    if len(ys) > 1:
        hints.append(f"Crossref dates span {ys}; using {rec.get('year')} (final version). Change with fix --year")
    if rec.get("container") and rec.get("type") != "book":
        missing = [f for f in ("volume", "pages") if not str(rec.get(f) or "").strip("0-")
                   and not (f == "pages" and rec.get("article_number"))]
        if missing:
            hints.append(f"no real {' / '.join(missing)} (in-press or incomplete record): take them from the journal "
                         f"page, then fix {rec['key']} --" + " --".join(f"{m} ..." for m in missing))
    if ARABIC_RE.search(t) and re.search(r"\s=\s*[A-Za-z]|\"[A-Za-z]", t):
        hints.append(f"title carries an English translation: fix {rec['key']} --title \"<Arabic title only>\"")
    return hints


def cmd_book(a):
    hits = book_lookup(isbn=a.isbn, title=a.title, author=a.author)
    if not hits:
        print("No match in Open Library. Do not cite the book until its ISBN is confirmed from the book itself.")
        return 1
    print(table([[str(i + 1), short(", ".join(x["family"] for x in h["authors"]), 30), str(h["year"]),
                  short(h["title"], 60), short("; ".join(h["publishers"]), 40), h.get("isbn") or "",
                  ",".join(map(str, h["publish_years"]))]
                 for i, h in enumerate(hits)],
                ["#", "authors", "first published", "title", "publishers", "isbn", "editions (years)"]))
    print("Cite the edition actually used: pass its year and publisher to `add --isbn ... --edition-year --publisher`.")


def cmd_abstract(a):
    doi = norm_doi(a.doi)
    msg = crossref_work(doi)
    rec = {"doi": doi, "retraction": {"status": "ok", "notices": []},
           "crossref_abstract": clean((msg or {}).get("abstract"))}
    rec = attach_abstract(rec)
    if rec.get("abstract"):
        print(f"[{rec['abstract_source']}] {rec['abstract']}")
    else:
        if S2_STATE["throttled"]:
            print("Semantic Scholar was rate-limited, so it was not really checked: retry in a minute, or set S2_API_KEY.")
        print("No abstract in OpenAlex, Crossref, Semantic Scholar or Europe PMC. Next: ask the user for the PDF, "
              "or open the open-access copy (Unpaywall: https://unpaywall.org, CORE: https://core.ac.uk). "
              "Until then cite it only for what its title shows; never summarise it from memory.")
        return 1


def cmd_cite(a):
    doi = norm_doi(a.doi)
    if a.style == "apa":
        rec, verdict = resolve_doi(doi)
        if rec is None:
            print(f"NOT CITABLE: {verdict}")
            return 1
        print(apa(rec))
    else:
        print(csl_formatted(doi, a.style, a.locale))


def cited_entries(ledger, manuscript):
    with open(manuscript, encoding="utf-8") as f:
        body, _ = strip_references(f.read())
    used = set()
    for sur, yr, _ in extract_citations(body):
        for e in ledger_match(sur, yr, ledger):
            used.add(e["key"])
    return used


def sort_key(e):
    """APA order: surname as written (Al-Zu'bi files under A), then initials,
    then year. Arabic-script names ignore the article ال, as Arab universities do."""
    fam = e["authors"][0]["family"] if e.get("authors") else e.get("title", "")
    if ARABIC_RE.search(fam):
        k = name_key(fam)
    else:
        k = unicodedata.normalize("NFKD", fam).lower()
        k = re.sub(r"[^a-z]", "", "".join(c for c in k if not unicodedata.combining(c)))
    given = e["authors"][0].get("given", "") if e.get("authors") else ""
    return (0 if ARABIC_RE.search(fam) else 1, k, initials(given), str(e.get("year")))


def cmd_refs(a):
    ledger = load_ledger(a.ledger)
    if a.manuscript:
        used = cited_entries(ledger, a.manuscript)
        ledger = [e for e in ledger if e["key"] in used]
    arabic_first = a.arabic_first
    ledger.sort(key=sort_key if arabic_first else (lambda e: sort_key(e)[1:]))
    for e in ledger:
        for h in record_hints(e):
            print(f"check {e['key']}: {h}", file=sys.stderr)
    if a.format != "apa":
        text = export_refs(ledger, a.format)
        if a.out:
            with open(a.out, "w", encoding="utf-8", newline="\n") as f:
                f.write(text)
            print(f"wrote {len(ledger)} {a.format} record(s) to {os.path.abspath(a.out)} (Zotero: File > Import)")
        else:
            sys.stdout.write(text)
        return 0
    out = [f"## {a.heading}\n"]
    for e in ledger:
        out += [e.get("apa") or apa(e), ""]
    if a.out:
        with open(a.out, "w", encoding="utf-8", newline="\n") as f:
            f.write("\n".join(out) + "\n")
        print(f"wrote {len(ledger)} reference(s) to {os.path.abspath(a.out)}")
    else:
        print("\n".join(out))


def cmd_fix(a):
    """Correct how a verified record is written, never what it is.

    Crossref sometimes splits Arabic names badly ('Samed Al-Adwan', 'Ahmad') and
    APA wants sentence-case titles. A new title must stay the same words.
    """
    ledger = load_ledger(a.ledger)
    e = next((x for x in ledger if x["key"] == a.key), None)
    if e is None:
        print(f"no ledger entry {a.key}")
        return 1
    # Validate everything first: a partly refused call changes nothing.
    changes, problems = [], []
    for spec in a.author:
        m = re.fullmatch(r"(\d+):(.+?),(.*)", spec)
        if not m or not (1 <= int(m.group(1)) <= len(e["authors"])):
            problems.append(f"bad --author '{spec}': use N:Family,Given with N between 1 and {len(e['authors'])}")
            continue
        i = int(m.group(1)) - 1
        old = e["authors"][i]
        new = {"family": m.group(2).strip(), "given": m.group(3).strip()}
        if not same_person(old, new):
            problems.append(f"'{spec}' is a different person from {old['family']}, {old['given']}: fix spelling, not identity")
            continue
        changes.append(("author", i, new))
    if a.title:
        if not same_title(e["title"], a.title):
            problems.append("--title may only change capitalisation, quotes, spacing or ®/™ symbols, or drop a "
                            "translation the registry appended ('Arabic title = English title'); not the words")
        else:
            changes.append(("title", None, clean(a.title)))
    for field in ("volume", "issue", "pages"):
        val = getattr(a, field)
        if val is None:
            continue
        old = str(e.get(field) or "")
        if old.strip("0-") and old != val:
            problems.append(f"--{field}: the record already has '{old}'. Only empty or 0 (in-press) values may be "
                            "filled, from the journal's own page")
        else:
            changes.append((field, None, clean(val)))
    if a.year:
        seen = e.get("years_seen") or []
        if e.get("doi") and not seen:
            msg = crossref_work(e["doi"]) or {}
            seen = crossref_years(msg)
        if a.year not in seen:
            problems.append(f"--year {a.year} is not a date Crossref holds for this work ({seen or 'none'})")
        else:
            changes.append(("year", None, a.year))
    if problems:
        for p in problems:
            print("REFUSED: " + p)
        print("Nothing was changed.")
        return 1
    for n_ in a.confirm_author:
        if 1 <= n_ <= len(e["authors"]):
            changes.append(("confirm", n_, None))
        else:
            print(f"REFUSED: --confirm-author {n_}: no such author position")
            return 1
    for kind, i, val in changes:
        if kind == "author":
            e["authors"][i] = val
        elif kind == "confirm":
            e.setdefault("confirmed_authors", [])
            if i not in e["confirmed_authors"]:
                e["confirmed_authors"].append(i)
        else:
            e[kind] = val
    e.setdefault("edits", []).append({"at": TODAY, "author": a.author, "title": bool(a.title), "year": a.year,
                                      "volume": a.volume, "issue": a.issue, "pages": a.pages})
    e["apa"] = apa(e)
    save_ledger(a.ledger, ledger)
    print(f"fixed {a.key}: {e['apa']}")


def same_title(old, new):
    sym = lambda s: re.sub(r"[®™©]", "", s)
    o, n = cmp_text(sym(old)).rstrip("."), cmp_text(sym(new)).rstrip(".")
    if o == n:
        return True
    # Registries append a translation: 'العنوان = Title' or 'العنوان "Title"'. Dropping it is allowed.
    return bool(n) and o.startswith(n) and re.match(r"""^[=/"'(\[:]""", o[len(n):] or "x") is not None


def same_person(old, new):
    """Same name tokens in any order and split (Crossref often puts the whole
    name in 'family' or swaps parts), allowing small spelling differences."""
    def toks(p):
        return [t for t in (name_key(x) for x in re.split(r"[\s,.\-]+", f"{p['family']} {p['given']}")) if len(t) > 1]
    o, n = toks(old), toks(new)
    if not o or not n:
        return False

    def near(t):
        return any(t == u or difflib.SequenceMatcher(None, t, u).ratio() >= 0.8 for u in o)

    # Reordered or re-split name: the new surname was already there and most tokens match.
    if near(name_key(new["family"])) and sum(map(near, n)) / len(n) >= 0.5:
        return True
    # Spelling or spacing fix ('Alameen' -> 'Al-Ameen').
    return difflib.SequenceMatcher(None, "".join(sorted(o)), "".join(sorted(n))).ratio() >= 0.85


def cmd_list(a):
    ledger = load_ledger(a.ledger)
    print(table([[e["key"], str(e.get("year")), short(e["authors"][0]["family"] if e.get("authors") else "", 18),
                  short(e.get("title"), 60),
                  e.get("doi") or (("ISBN " + str(e["isbn"])) if e.get("isbn") else "manual (unverified)"),
                  (e.get("retraction") or {}).get("status", ""), e.get("abstract_source") or "none",
                  e.get("read_level", ""), short(e.get("note"), 40)] for e in ledger],
                ["key", "year", "first author", "title", "id", "integrity", "abstract", "read", "supports"]))


def cmd_audit(a):
    a.manuscript = a.manuscript or a.manuscript_opt
    if not a.manuscript:
        sys.exit("audit needs the manuscript file: audit paper.md (or --manuscript paper.md)")
    ledger = load_ledger(a.ledger)
    with open(a.manuscript, encoding="utf-8") as f:
        text = f.read()
    body, refsec = strip_references(text)
    cites = extract_citations(body)
    fails, warns = [], []
    used = set()
    if a.keys:                      # coaching notes and matrices name sources by ledger key
        for e in ledger:
            if re.search(r"(?<![\w])" + re.escape(e["key"]) + r"(?![\w])", body):
                used.add(e["key"])
    for sur, yr, snip in cites:
        hits = ledger_match(sur, yr, ledger)
        if not hits:
            fails.append(f"NOT IN LEDGER: {sur} ({yr}) in \"{snip}\"")
            continue
        if len(hits) > 1 and not re.search(r"[a-z]$", yr):
            warns.append(f"AMBIGUOUS: {sur} ({yr}) matches {', '.join(h['key'] for h in hits)}: add a/b suffixes")
        for h in hits:
            used.add(h["key"])
            st = (h.get("retraction") or {}).get("status")
            if st in ("retracted", "expression_of_concern"):
                fails.append(f"{st.upper()}: {h['key']} cited in \"{snip}\"")
            if not h.get("abstract") and h.get("read_level") not in ("full", "full_text"):
                warns.append(f"UNREAD: {h['key']} has no abstract and was not marked as read in full")
    for e in ledger:
        if e["key"] not in used and not a.partial:
            warns.append(f"UNCITED: {e['key']} is in the ledger but not cited in the text")
        if e["key"] in used:
            warns += [f"{e['key']}: {h}" for h in record_hints(e) if h.startswith("Title Case")]
    if refsec.strip():
        ledger_dois = {(e.get("doi") or "").lower() for e in ledger}
        for line in [l.strip() for l in refsec.splitlines() if l.strip()]:
            dois = [norm_doi(d).lower() for d in re.findall(r"10\.\d{4,9}/[^\s<>\"]+", line)]
            if dois:
                for d in dois:
                    if d.rstrip(".") not in ledger_dois:
                        fails.append(f"REFERENCE NOT IN LEDGER: {short(line, 100)}")
            else:
                best = max((title_sim(e.get("title", ""), line) for e in ledger), default=0)
                inside = any(clean(e.get("title", "")).lower()[:40] in clean(line).lower() for e in ledger if e.get("title"))
                if not inside and best < 0.5:
                    fails.append(f"REFERENCE NOT IN LEDGER (no DOI): {short(line, 100)}")
    elif not a.partial:
        warns.append("REFERENCE SECTION missing or empty (heading 'References' or 'المراجع'): build it with `refs`")
    if any(ch in text for ch in EN_EM):
        n = sum(text.count(ch) for ch in EN_EM)
        warns.append(f"DASHES: {n} en/em dashes found; replace with hyphen, comma, colon or parentheses")
    if not a.partial and not re.search(r"(generative ai|ai-assisted|artificial intelligence|الذكاء الاصطناعي)", text, re.I):
        warns.append("NO AI-USE DISCLOSURE found: add the statement your target journal or university requires")
    print(f"manuscript: {a.manuscript} | citations found: {len(cites)} | ledger: {len(ledger)} sources")
    if not cites and not used:
        print("NOTE  no in-text citations found: nothing was checked against the ledger. "
              "For a bare reference list use `verify --citations-file`.")
    for x in fails:
        print("FAIL  " + x)
    for x in dict.fromkeys(warns):
        print("WARN  " + x)
    print("RESULT: " + ("FAIL" if fails else "PASS") + f" ({len(fails)} fail, {len(set(warns))} warn)")
    return 1 if fails else 0


def cmd_numbers(a):
    with open(a.text, encoding="utf-8") as f:
        items = numbers_in_text(f.read(), include_ints=bool(a.ledger))
    show = lambda i: f"line {i['line']}: {i['label'] or '?'} {i['op']}{i['value']}{'%' if i['pct'] else ''}  | {i['context']}"
    if not a.evidence and not a.ledger:
        # Template mode: before the data exists, a results text holds placeholders only.
        print(f"text: {a.text} | template mode (no evidence given): every statistic must be a [ ] placeholder")
        for i in items:
            print("STATISTIC IN TEMPLATE " + show(i))
        print("RESULT: " + ("FAIL" if items else "PASS") + f" ({len(items)} statistics found)")
        return 1 if items else 0
    texts = []
    if a.ledger:
        texts = [e.get("abstract", "") for e in load_ledger(a.ledger) if e.get("abstract")]
    ev = evidence_numbers(a.evidence or [], texts)
    res = [(i, number_supported(i, ev)) for i in items]
    missing = [i for i, r in res if r is None]
    sign = [i for i, r in res if r == "sign"]
    weak = [i for i, r in res if r == "weak"]
    src = f"{len(a.evidence or [])} file(s)" + (f" + {len(texts)} ledger abstracts" if a.ledger else "")
    print(f"text: {a.text} | numbers checked: {len(items)} | evidence values: {len(ev)} from {src}")
    for i in missing:
        print("UNSUPPORTED " + show(i))
    for i in sign:
        print("SIGN DIFFERS " + show(i))
    for i in weak:
        print("WEAK MATCH  " + show(i))
    if weak and a.ledger and not a.evidence:
        print(f"  {len(weak)} weak match(es) are expected in ledger mode (abstracts are prose without labels): "
              "confirm each value appears in the abstract of the source cited next to it.")
    elif weak:
        print(f"  {len(weak)} value(s) matched only an unlabelled number. Write the analysis output as "
              "'label = value' lines (scripts/analyze.py does; the LAST segment of the key must be the "
              "statistic, e.g. 'mytest.d = 0.40') so each statistic is checked against its own label.")
    fail = missing + sign
    print("RESULT: " + ("FAIL" if fail else "PASS") + f" ({len(missing)} unsupported, {len(sign)} sign, {len(weak)} weak)")
    return 1 if fail else 0


LISTED_IN = {
    "abdc": "ABDC (Australian business journal list)", "erih-plus": "ERIH PLUS (European humanities/social sciences)",
    "jufo": "JUFO (Finnish Publication Forum level)", "norway": "Norwegian Register level",
    "anvur": "ANVUR (Italian evaluation list)", "fnege": "FNEGE (French management list)",
    "cwts-core": "CWTS core international journal", "doaj": "DOAJ", "medline": "MEDLINE",
    "vabb": "VABB (Flemish academic bibliography)", "poland": "Polish ministry points list",
    "russia-white-list": "Russian white list",
}


PUBLISHER_FAMILIES = {
    "springer": ("springer", "kluwer", "nature", "biomed central", "bmc", "palgrave"),
    "elsevier": ("elsevier", "cell press", "academic press", "pergamon", "north-holland"),
    "taylor": ("taylor", "routledge", "informa", "crc press", "dove medical"),
    "wiley": ("wiley", "blackwell", "hindawi"),
    "sage": ("sage",), "emerald": ("emerald",), "ieee": ("ieee", "institute of electrical"),
    "oup": ("oxford university press",), "cup": ("cambridge university press",), "mdpi": ("mdpi",),
    "frontiers": ("frontiers",), "acm": ("association for computing", "acm"),
}


def publisher_family(name):
    """Imprints and merged houses count as one publisher (Kluwer = Springer, Routledge = T&F)."""
    n = clean(name).lower()
    for fam, words in PUBLISHER_FAMILIES.items():
        if any(w in n for w in words):
            return fam
    return re.sub(r"[^a-z]", "", n)[:12]


def journal_key(title):
    t = clean(title).lower().replace("&", "and")
    t = re.sub(r"^the\s+", "", t)
    return re.sub(r"[^a-z0-9؀-ۿ]", "", t)


def edition_warning(edition):
    try:
        if int(edition) < dt.date.today().year - 1:
            print(f"Note: this is the SJR {edition} edition; a newer edition is probably out. "
                  "Re-download from scimagojr.com before quoting quartiles.")
    except ValueError:
        pass


def explain_listed(codes):
    out = []
    for c in codes:
        c = str(c)
        base = next((k for k in LISTED_IN if c == k or c.startswith(k + "-") or c.startswith(k)), None)
        out.append(f"{c} = {LISTED_IN[base]}" if base else c)
    return "; ".join(out)


def cmd_journal(a):
    for n, ident in enumerate(a.ident):
        if n:
            print()
            print("=" * 70)
        journal_one(a, ident)


def journal_one(a, ident):
    issn = hyphen_issn(ident)
    print(f"journal: {ident} | checked {TODAY}")
    sjr_path = a.sjr or os.environ.get("SJR_CSV")
    sjr_hit = None
    if sjr_path:
        journals, edition = load_sjr(sjr_path)
        ends = [int(j["coverage"][-4:]) for j in journals if re.search(r"\d{4}$", j["coverage"])]
        max_cov = max(set(ends), key=ends.count) if ends else 0      # the usual end year of active titles
        if issn:
            hits = [j for j in journals if compact_issn(issn) in j["issns"]]
        else:
            # A title must match exactly (ignoring case, '&'/'and', punctuation): a near
            # miss is a different journal, and fake journals borrow real-sounding names.
            want = journal_key(ident)
            hits = [j for j in journals if journal_key(j["title"]) == want]
            if not hits:
                near = sorted(journals, key=lambda j: -title_sim(j["title"], ident))[:3]
                near = [j for j in near if title_sim(j["title"], ident) > 0.6]
                if near:
                    print(f"\nSJR {edition}: no journal with exactly this title. Similar real titles (NOT the same journal): "
                          + "; ".join(f"{j['title']} ({', '.join(hyphen_issn(x) for x in j['issns'])})" for j in near))
        if hits:
            sjr_hit = hits[0]
            for j in hits:
                print(f"\nSJR {edition} (SCImago, file supplied by user): {j['title']} | ISSN "
                      f"{', '.join(hyphen_issn(x) for x in j['issns'])}")
                print(f"  SJR {j['sjr']} | best quartile {j['best_q']} | H index {j['h']} | {j['type']} | "
                      f"OA {j['oa']} | {j['publisher']} | {j['country']} ({j['region']}) | coverage {j['coverage']}")
                for n, q in j["categories"]:
                    print(f"  - {n}: {q}")
                print("  Quote the quartile of the paper's own category, not only the best quartile.")
                end = re.search(r"(\d{4})$", j["coverage"])
                if end and max_cov and int(end.group(1)) < max_cov:
                    print(f"  WARNING: Scopus coverage ends {end.group(1)} while active titles run to {max_cov}: "
                          "possible coverage gap or discontinuation. Check the Scopus Discontinued Titles list.")
            if not issn:
                issn = hyphen_issn(sjr_hit["issns"][0]) if sjr_hit["issns"] else ""
        else:
            print(f"\nSJR {edition}: NOT FOUND. Not indexed in Scopus, or the ISSN/title differs: treat 'Q1' claims as unproven.")
        edition_warning(edition)
    else:
        print("\nSJR: no file supplied. Quartiles cannot be confirmed without the user's SCImago download "
              "(see references/journal-vetting.md). Do not state a quartile.")
    src = None
    if issn:
        try:
            src = openalex("sources/issn:" + issn)
        except HTTPError as e:
            src = None
            print(f"\nOpenAlex: unavailable for this ISSN (HTTP {e.status}).")
    elif ident.strip():
        data = openalex("sources", {"search": ident, "per_page": 3,
                                    "select": "id,display_name,issn_l,issn,host_organization_name,works_count,"
                                              "summary_stats,is_in_doaj,apc_usd,country_code,homepage_url,listed_in"})
        res = (data or {}).get("results", [])
        cands = [s for s in res if journal_key(s.get("display_name", "")) == journal_key(ident)]
        src = cands[0] if cands else None
        if not src and res:
            print("\nOpenAlex: no source with exactly this title. Closest (NOT the same journal): "
                  + "; ".join(f"{s.get('display_name')} ({s.get('issn_l')})" for s in res[:3]))
        if src and src.get("issn_l"):
            issn = src["issn_l"]
    if src:
        st = src.get("summary_stats") or {}
        print(f"\nOpenAlex: {src.get('display_name')} | ISSN-L {src.get('issn_l')} | {src.get('host_organization_name')} "
              f"| works {src.get('works_count')} | h-index {st.get('h_index')} "
              f"| 2yr mean citedness {round(st.get('2yr_mean_citedness') or 0, 2)} | in DOAJ {src.get('is_in_doaj')} "
              f"| {src.get('country_code')} | homepage {src.get('homepage_url')}")
        if src.get("apc_usd"):
            print(f"  APC (index list price, may be stale): USD {src['apc_usd']}. Confirm on the journal's own site.")
        if src.get("listed_in"):
            print(f"  listed in: {explain_listed(src['listed_in'])[:700]}")
        if sjr_hit and src.get("host_organization_name") and \
                publisher_family(sjr_hit["publisher"]) != publisher_family(src["host_organization_name"]):
            print(f"  CHECK: publishers disagree (SJR '{sjr_hit['publisher']}' vs OpenAlex "
                  f"'{src['host_organization_name']}'): confirm on the journal site.")
    else:
        print(f"\nOpenAlex: no source found for '{ident}'. An unknown journal name is itself a warning sign.")
    if issn:
        try:
            r = http_get("https://doaj.org/api/search/journals/issn:" + issn)
            res = as_json(r).get("results", []) if r is not None else []
        except HTTPError as e:
            res = None
            print(f"\nDOAJ: unavailable (HTTP {e.status}); OpenAlex says in DOAJ = "
                  f"{src.get('is_in_doaj') if src else 'unknown'} (fallback signal only).")
        if res:
            bj = res[0].get("bibjson", {})
            apc = bj.get("apc", {})
            fee = ", ".join(f"{p.get('price')} {p.get('currency')}" for p in apc.get("max", [])) if apc.get("has_apc") else "none"
            print(f"\nDOAJ: listed | APC {fee} (DOAJ record, may be stale: confirm on the journal site) "
                  f"| review {', '.join((bj.get('editorial') or {}).get('review_process', []))} "
                  f"| licence {', '.join(l.get('type', '') for l in bj.get('license', []))} | site {(bj.get('ref') or {}).get('journal')}")
        elif res is not None:
            print("\nDOAJ: not listed (normal for subscription journals; a red flag only if it claims to be open access).")
    if a.hijacked:
        check_hijacked(a.hijacked, issn, a.url)
    elif a.url:
        print("\nHijack check skipped: pass --hijacked <Retraction Watch checker .xlsx> to compare the website.")
    print("\nThese are signals, not a verdict. Finish with the Think.Check.Submit checklist "
          "(Arabic: https://thinkchecksubmit.org/journals/arabic/).")


def check_hijacked(path, issn, url):
    try:
        import openpyxl
    except ImportError:
        print("openpyxl missing: pip install openpyxl")
        return
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    stamp = next((str(c) for r in rows[:3] for c in (r or []) if c and "updated" in str(c).lower()), "date unknown")
    hi = next(i for i, r in enumerate(rows) if r and any(str(c or "").strip().lower().startswith("url") for c in r))
    head = [str(c or "").strip().lower() for c in rows[hi]]
    col = lambda word: next((i for i, h in enumerate(head) if word in h), None)
    c_hurl, c_hissn, c_ourl = col("url (hijacked)"), col("issn (hijacked)"), col("url (original")
    data = [r for r in rows[hi + 1:] if r and any(r)]
    print(f"\nRetraction Watch hijacked-journal list ({stamp}; {len(data)} entries; file {os.path.basename(path)})")

    def dom_of(u):
        u = str(u or "").strip()
        return urllib.parse.urlparse(u if u.startswith("http") else "http://" + u).netloc.lower().removeprefix("www.")

    def site_of(u):
        """Domain + path, so journals sharing a platform (sciencedirect.com, link.springer.com) stay apart."""
        u = str(u or "").strip()
        pu = urllib.parse.urlparse(u if u.startswith("http") else "http://" + u)
        return pu.netloc.lower().removeprefix("www.") + pu.path.rstrip("/").lower()

    def same_site(a_, b_):
        if not a_ or not b_:
            return False
        if a_ == b_:
            return True
        # A listed bare domain ('aomannals.com') covers every page on it; a listed
        # platform path only covers that journal's own pages.
        short_, long_ = sorted((a_, b_), key=len)
        return long_.startswith(short_ + "/") and ("/" in short_ or short_ == dom_of(long_))

    given = site_of(url) if url else ""
    issn_hits, url_hits, official = [], [], []
    for r in data:
        cell = lambda i: str(r[i] or "") if i is not None and i < len(r) else ""
        if issn and compact_issn(issn) in re.sub(r"[^0-9Xx,]", "", cell(c_hissn)).upper().split(","):
            issn_hits.append(r)
        if given:
            if same_site(given, site_of(cell(c_hurl))):
                url_hits.append(r)
            elif same_site(given, site_of(cell(c_ourl))):
                official.append(r)
    if url_hits:
        print(f"  DANGER: {url} IS a known hijacked site. Official site: "
              f"{url_hits[0][c_ourl] if c_ourl is not None else 'see the list'}")
    if official:
        print(f"  {url} matches the OFFICIAL site listed for this journal (a clone exists: never use the clone).")
    if issn_hits and not url_hits:
        h = issn_hits[0]
        print(f"  This ISSN has a known hijacked clone at {h[c_hurl] if c_hurl is not None else '?'}; "
              f"official site: {h[c_ourl] if c_ourl is not None else '?'}. Submit only through the official site.")
    if not (issn_hits or url_hits or official):
        print("  No match for this ISSN or domain.")
    if issn and not url:
        print("  Pass --url <the site you were given> to check that exact website.")


def cmd_journals(a):
    journals, edition = load_sjr(a.sjr or os.environ.get("SJR_CSV"))
    out = []
    for j in journals:
        if a.type and j["type"] != a.type:
            continue
        if a.oa and j["oa"].lower() != a.oa:
            continue
        if a.region and a.region.lower() not in j["region"].lower():
            continue
        wanted = [c.strip().lower() for c in a.category.split("|")]   # '|' separates; commas belong to names
        if a.exact:
            cq = [(n, q) for n, q in j["categories"] if n.lower() in wanted]
        else:
            cq = [(n, q) for n, q in j["categories"] if any(w in n.lower() for w in wanted)]
        if not cq:
            continue
        q = min((q for _, q in cq if q.startswith("Q")), default="-")
        if a.quartile and q not in a.quartile.split(","):
            continue
        out.append((j, "; ".join(f"{n} ({q2})" for n, q2 in cq), q))
    out.sort(key=lambda t: -(t[0]["sjr"] or 0))
    mode = "is" if a.exact else "contains"
    print(f"SJR {edition} | category {mode} '{a.category}' | {len(out)} journals | showing {min(len(out), a.limit)}")
    print(table([[str(i + 1), short(j["title"], 70), ", ".join(hyphen_issn(s) for s in j["issns"]), str(j["sjr"]),
                  q, short(n, 70), j["h"], j["oa"], short(j["publisher"], 30), j["country"]]
                 for i, (j, n, q) in enumerate(out[: a.limit])],
                ["#", "journal", "issn", "SJR", "best quartile here", "matching categories", "H", "OA", "publisher", "country"]))
    if not a.exact:
        names = sorted({n for j, _, _ in out for n, _ in j["categories"]
                        if any(w in n.lower() for w in [c.strip().lower() for c in a.category.split("|")])})
        print(f"Categories matched: {'; '.join(names)[:600]}")
        print("Too broad? Rerun with --exact and the exact category name(s), comma-separated.")
    edition_warning(edition)
    print(f"Source: SCImago (n.d.). SJR SCImago Journal & Country Rank [Portal]. Retrieved {TODAY}, from https://www.scimagojr.com")


# ---------------------------------------------------------------- cli

def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)

    s = sub.add_parser("search", help="find candidate papers")
    s.add_argument("query", nargs="?", default="", help="quoted phrases with AND/OR work best; optional with --issn")
    s.add_argument("--issn", help="only papers in these journals (comma-separated ISSNs): scope checks")
    s.add_argument("--core", action="store_true", help="only journals in the CWTS core list (drops most local/predatory outlets)")
    s.add_argument("--has-doi", action="store_true", help="only works with a DOI (citable through `add`)")
    s.add_argument("--source", choices=["openalex", "s2", "both", "crossref"], default="openalex",
                   help="crossref covers Arabic journals OpenAlex misses (EKB 10.21608, Al Manhal 10.12816, "
                        "Jordanian and Gulf university journals). Example: search \"القيادة التحويلية الأداء الوظيفي\" --source crossref")
    s.add_argument("--expand", action="store_true",
                   help="also run the Arabic synonym variants of the query (see --show-expansions) and merge the hits")
    s.add_argument("--show-expansions", action="store_true", help="print the built-in Arabic synonym map and exit")
    s.add_argument("--since", type=int)
    s.add_argument("--until", type=int)
    s.add_argument("--type", default="article", help="OpenAlex work type (article, review, book-chapter...); '' for any")
    s.add_argument("--min-cites", type=int, default=0)
    s.add_argument("--oa", action="store_true", help="open access only")
    s.add_argument("--lang", help="ISO 639-1 language, e.g. ar or en")
    s.add_argument("--sort", choices=["relevance", "cites", "recent"], default="relevance")
    s.add_argument("--limit", type=int, default=25)
    s.add_argument("--abstracts", action="store_true", help="include abstracts in --out")
    s.add_argument("--out", help="save candidates as JSON")

    v = sub.add_parser("verify", help="check DOIs exist + retraction status, or match a citation string")
    v.add_argument("dois", nargs="*")
    v.add_argument("--citation", help="free-text reference to match to a DOI")
    v.add_argument("--citations-file", help="a whole reference list (one per line or per blank-line block)")
    v.add_argument("--add-to", help="with --citations-file: add the confirmed ones to this ledger")

    ad = sub.add_parser("add", help="verify and append sources to the ledger")
    ad.add_argument("dois", nargs="*")
    ad.add_argument("--isbn", action="append", default=[])
    ad.add_argument("--ledger", default="sources.jsonl")
    ad.add_argument("--note", help="what claim this source supports")
    ad.add_argument("--read", choices=["abstract", "full", "metadata"], default="abstract",
                    help="what has actually been read")
    ad.add_argument("--allow-retracted", action="store_true")
    ad.add_argument("--edition-year", type=int)
    ad.add_argument("--publisher")
    ad.add_argument("--dois-file", help="a text file with one DOI per line (batch add)")
    ad.add_argument("--manual", action="store_true",
                    help="a source with no DOI or ISBN (scale manual, report, thesis): needs --title, --year, "
                         "--authors \"Family, Given; ...\"; stored as unverified")
    ad.add_argument("--title")
    ad.add_argument("--authors", help='with --manual: "Family, Given; Family, Given"')
    ad.add_argument("--year", type=int)
    ad.add_argument("--type", default="report", help="with --manual: report, manual, thesis, standard, web")
    ad.add_argument("--url")

    b = sub.add_parser("book", help="look up a book (Open Library)")
    b.add_argument("--isbn")
    b.add_argument("--title")
    b.add_argument("--author")

    ab = sub.add_parser("abstract", help="real abstract of a DOI")
    ab.add_argument("doi")

    c = sub.add_parser("cite", help="formatted reference")
    c.add_argument("doi")
    c.add_argument("--style", default="apa", help="apa (built locally) or any CSL style name via doi.org, e.g. ieee, harvard-cite-them-right")
    c.add_argument("--locale", help="CSL locale, e.g. ar")

    r = sub.add_parser("refs", help="reference list from the ledger")
    r.add_argument("--ledger", default="sources.jsonl")
    r.add_argument("--manuscript", help="only include sources this file cites")
    r.add_argument("--heading", default="References")
    r.add_argument("--arabic-first", action="store_true", help="Arabic-script references first (common in Arab universities)")
    r.add_argument("--format", choices=["apa", "ris", "bibtex", "csl"], default="apa",
                   help="apa (markdown list, default) or a Zotero import format: ris, bibtex (keys = ledger keys), "
                        "csl (CSL-JSON). Example: refs --manuscript paper.md --format ris --out refs.ris")
    r.add_argument("--out", help="write to this file instead of stdout (UTF-8)")

    fx = sub.add_parser("fix", help="correct author spelling or title case of a ledger entry")
    fx.add_argument("key")
    fx.add_argument("--author", action="append", default=[], help="N:Family,Given (1-based position)")
    fx.add_argument("--title", help="same title in APA sentence case")
    fx.add_argument("--year", type=int, help="final-version year; must be a date Crossref holds for the work")
    fx.add_argument("--volume", help="fill a missing / 0 volume (in-press records) from the journal page")
    fx.add_argument("--issue", help="fill a missing / 0 issue")
    fx.add_argument("--pages", help="fill missing / 0-0 pages")
    fx.add_argument("--confirm-author", action="append", type=int, default=[],
                    help="N: this author's name is right as stored (silences the 'may be split wrongly' hint)")
    fx.add_argument("--ledger", default="sources.jsonl")

    li = sub.add_parser("list", help="show the ledger")
    li.add_argument("--ledger", default="sources.jsonl")

    au = sub.add_parser("audit", help="check manuscript citations against the ledger")
    au.add_argument("manuscript", nargs="?")
    au.add_argument("--manuscript", dest="manuscript_opt", help="same as the positional argument")
    au.add_argument("--ledger", default="sources.jsonl")
    au.add_argument("--partial", action="store_true",
                    help="one section or chapter: skip the uncited, reference-section and disclosure warnings")
    au.add_argument("--keys", action="store_true", help="also count ledger keys written in the text as citations")

    n = sub.add_parser("numbers", help="every statistic in the text must appear in analysis output")
    n.add_argument("text")
    n.add_argument("--evidence", nargs="+", help="analysis output files (txt, csv, json); omit both this and "
                   "--ledger to check a results TEMPLATE has no statistics")
    n.add_argument("--ledger", help="use the ledger abstracts as evidence (numbers quoted in a literature review)")

    j = sub.add_parser("journal", help="vet one journal by ISSN or title")
    j.add_argument("ident", nargs="+", help="one or more ISSNs or exact titles")
    j.add_argument("--sjr", help="SCImago CSV downloaded by the user")
    j.add_argument("--hijacked", help="Retraction Watch hijacked journal checker .xlsx")
    j.add_argument("--url", help="the journal website you were given, for the hijack check")

    js = sub.add_parser("journals", help="shortlist journals from an SJR CSV")
    js.add_argument("--sjr")
    js.add_argument("--category", required=True,
                    help="SJR category text, e.g. 'Education'; several with '|' ('Education|E-learning'), since "
                         "category names contain commas ('Public Health, Environmental and Occupational Health')")
    js.add_argument("--exact", action="store_true", help="match whole category names only")
    js.add_argument("--quartile", help="e.g. Q1 or Q1,Q2")
    js.add_argument("--type", default="journal")
    js.add_argument("--oa", choices=["yes", "no"])
    js.add_argument("--region", help="e.g. 'Middle East'")
    js.add_argument("--limit", type=int, default=25)

    a = p.parse_args()
    if a.cmd == "search" and a.type == "":
        a.type = None
    if a.cmd == "verify" and not (a.dois or a.citation or a.citations_file):
        p.error("verify needs DOIs, --citation or --citations-file")
    if a.cmd == "search" and not (a.query or a.issn or a.show_expansions):
        p.error("search needs a query or --issn")
    if a.cmd == "book" and not (a.isbn or a.title):
        p.error("book needs --isbn or --title")
    fn = globals()["cmd_" + a.cmd]
    try:
        sys.exit(fn(a) or 0)
    except HTTPError as e:
        sys.exit(f"API error: {e} {e.body[:200] if isinstance(e.body, str) else ''}")


if __name__ == "__main__":
    main()
